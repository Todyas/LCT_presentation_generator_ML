from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

from app.config import Settings
from app.pipeline import jobs as jobs_module
from app.pipeline.orchestrator import Dependencies
from tests.fakes import ScriptedLLM

REPO_ROOT = Path(__file__).resolve().parents[1]
SKILLS_DIR = REPO_ROOT / "skills"
FIXTURE_TEMPLATE = (
    REPO_ROOT / "tests" / "fixtures" / "templates" / "generated_minimal.pptx"
)


@pytest.fixture
def tmp_cwd(tmp_path, monkeypatch):
    """Chdir into an isolated tmp_path so TemplateParser's `.cache/` and any
    accidental `.env` reads never touch the real repository."""
    monkeypatch.chdir(tmp_path)
    return tmp_path


@pytest.fixture
def test_settings(tmp_path) -> Settings:
    # `:memory:` sqlite engines are cached process-wide by URL (see
    # app.pipeline.jobs.get_engine's _ENGINES dict). Dropping the cache entry
    # here forces the *next* JobStore/SlideRevisionStore construction in this
    # test to create a brand-new StaticPool, so tests never see each other's
    # rows even though they share the same database_url string.
    jobs_module._ENGINES.pop("sqlite+pysqlite:///:memory:", None)
    return Settings(
        _env_file=None,
        database_url="sqlite+pysqlite:///:memory:",
        task_queue_enabled=False,
        storage_dir=str(tmp_path / "storage"),
        skills_dir=str(SKILLS_DIR),
    )


@pytest.fixture
def scripted_llm() -> ScriptedLLM:
    return ScriptedLLM()


@pytest.fixture
def deps(test_settings, scripted_llm) -> Dependencies:
    dependencies = Dependencies(test_settings)
    dependencies.llm = scripted_llm
    return dependencies


@pytest.fixture
def patched_exporters(tmp_path):
    """Patch the three export functions at their orchestrator import sites.

    Default behaviour: pdf export "succeeds" (returns a path that need not
    exist), no previews are produced, and the html exporter is therefore
    never invoked by the orchestrator's own `if preview_paths:` guard.
    """
    pdf_path = str(tmp_path / "built.pdf")
    convert_mock = mock.AsyncMock(return_value=pdf_path)
    previews_mock = mock.Mock(return_value=[])
    html_mock = mock.Mock(return_value=str(tmp_path / "viewer.html"))
    with (
        mock.patch("app.pipeline.orchestrator.convert_to_pdf", new=convert_mock),
        mock.patch("app.pipeline.orchestrator.render_previews", new=previews_mock),
        mock.patch("app.pipeline.orchestrator.export_html_viewer", new=html_mock),
    ):
        yield SimpleNamespace(
            convert_to_pdf=convert_mock,
            render_previews=previews_mock,
            export_html_viewer=html_mock,
            pdf_path=pdf_path,
        )


@pytest.fixture
def patch_exporters_with_previews(tmp_path):
    """Factory fixture: call `apply(n)` to patch the exporters so that
    `render_previews` returns `n` real PNG files written under tmp_path."""

    def _apply(n: int, *, pdf_raises: Exception | None = None):
        pdf_path = str(tmp_path / "built.pdf")
        if pdf_raises is not None:
            convert_mock = mock.AsyncMock(side_effect=pdf_raises)
        else:
            convert_mock = mock.AsyncMock(return_value=pdf_path)
        preview_paths = []
        for i in range(n):
            png = tmp_path / f"preview-{i + 1:02d}.png"
            png.write_bytes(b"\x89PNG\r\n\x1a\n")
            preview_paths.append(str(png))
        previews_mock = mock.Mock(return_value=preview_paths)
        html_path = str(tmp_path / "viewer.html") if preview_paths else None
        html_mock = mock.Mock(return_value=html_path)
        patchers = [
            mock.patch("app.pipeline.orchestrator.convert_to_pdf", new=convert_mock),
            mock.patch("app.pipeline.orchestrator.render_previews", new=previews_mock),
            mock.patch("app.pipeline.orchestrator.export_html_viewer", new=html_mock),
        ]
        for p in patchers:
            p.start()
        return SimpleNamespace(
            convert_to_pdf=convert_mock,
            render_previews=previews_mock,
            export_html_viewer=html_mock,
            preview_paths=preview_paths,
            patchers=patchers,
        )

    applied: list = []
    original_apply = _apply

    def _tracking_apply(*args, **kwargs):
        result = original_apply(*args, **kwargs)
        applied.append(result)
        return result

    yield _tracking_apply
    for result in applied:
        for p in result.patchers:
            p.stop()


# --------------------------------------------------------------------------- #
# Synthetic template files
# --------------------------------------------------------------------------- #


@pytest.fixture(scope="session")
def default_template(tmp_path_factory) -> str:
    """A stock python-pptx `Presentation()`: 11 standard layouts with placeholders."""
    out_dir = tmp_path_factory.mktemp("default_template")
    path = out_dir / "default.pptx"
    Presentation().save(str(path))
    return str(path)


@pytest.fixture(scope="session")
def placeholderless_template(tmp_path_factory) -> str:
    """Two example slides on the Blank layout using only add_textbox shapes,
    to exercise `_inferred_slots` mining of example content."""
    out_dir = tmp_path_factory.mktemp("placeholderless_template")
    path = out_dir / "placeholderless.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for i in range(2):
        slide = prs.slides.add_slide(blank)
        title_box = slide.shapes.add_textbox(
            Emu(int(prs.slide_width * 0.08)),
            Emu(int(prs.slide_height * 0.06)),
            Emu(int(prs.slide_width * 0.84)),
            Emu(int(prs.slide_height * 0.14)),
        )
        title_box.name = f"Example Title {i}"
        title_box.text_frame.text = f"Example heading {i}"
        body_box = slide.shapes.add_textbox(
            Emu(int(prs.slide_width * 0.08)),
            Emu(int(prs.slide_height * 0.26)),
            Emu(int(prs.slide_width * 0.84)),
            Emu(int(prs.slide_height * 0.6)),
        )
        body_box.name = f"Example Body {i}"
        body_box.text_frame.text = f"Example body content {i}"
    prs.save(str(path))
    return str(path)


@pytest.fixture(scope="session")
def asymmetric_template(tmp_path_factory) -> str:
    """An example slide on Blank with a "brand artwork" picture covering the
    right 50% and a textbox confined to the left 45%.

    The artwork is a real Picture shape (not a filled AutoShape) on purpose:
    an AutoShape always has a `text_frame` — even an empty one — so
    TemplateParser._inferred_slots would happily mine it as a second,
    ambiguous BODY content slot alongside the real textbox. A Picture has no
    text_frame at all, so it is correctly never mistaken for a content slot.
    """
    out_dir = tmp_path_factory.mktemp("asymmetric_template")
    path = out_dir / "asymmetric.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    import io

    from PIL import Image

    image_bytes = io.BytesIO()
    Image.new("RGB", (4, 4), (0x11, 0x22, 0x99)).save(image_bytes, format="PNG")
    image_bytes.seek(0)
    artwork = slide.shapes.add_picture(
        image_bytes,
        Emu(int(prs.slide_width * 0.5)),
        Emu(0),
        Emu(int(prs.slide_width * 0.5)),
        Emu(prs.slide_height),
    )
    artwork.name = "brand artwork"

    body_box = slide.shapes.add_textbox(
        Emu(int(prs.slide_width * 0.05)),
        Emu(int(prs.slide_height * 0.25)),
        Emu(int(prs.slide_width * 0.40)),
        Emu(int(prs.slide_height * 0.5)),
    )
    body_box.name = "Example Body"
    body_box.text_frame.text = "Example left-column content"
    prs.save(str(path))
    return str(path)


@pytest.fixture(scope="session")
def dark_template(tmp_path_factory) -> str:
    """A near-black master background, to exercise contrast logic."""
    out_dir = tmp_path_factory.mktemp("dark_template")
    path = out_dir / "dark.pptx"
    prs = Presentation()
    master = prs.slide_masters[0]
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
    prs.save(str(path))
    return str(path)


@pytest.fixture
def generated_minimal_template() -> str:
    return str(FIXTURE_TEMPLATE)
