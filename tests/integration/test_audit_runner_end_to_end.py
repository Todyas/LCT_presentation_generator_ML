import hashlib
import shutil
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from pptx import Presentation

from app.core.agents.prompt_registry import PromptRegistry
from app.core.auditor.audit_runner import run_full_audit
from app.core.auditor.semantic_audit import SemanticFindings
from app.core.builder.pptx_builder import PptxBuilder
from app.core.parser.template_parser import TemplateParser
from app.models.audit_report import IssueType
from app.models.presentation_ir import (
    BulletBlock,
    BulletItem,
    PresentationIR,
    SlideIR,
    TitleComponent,
)
from app.models.template_manifest import LayoutType, PlaceholderType

FIXTURE_PATH = "tests/fixtures/templates/generated_minimal.pptx"


def _cache_path_for(pptx_path: str) -> Path:
    file_bytes = Path(pptx_path).read_bytes()
    source_hash = hashlib.sha256(file_bytes).hexdigest()
    return Path(".cache") / f"{source_hash}.manifest.json"


@pytest.fixture(autouse=True)
def _clear_cache():
    cache_path = _cache_path_for(FIXTURE_PATH)
    cache_path.unlink(missing_ok=True)
    yield
    cache_path.unlink(missing_ok=True)


def _copy_template(tmp_path: Path) -> str:
    dest = tmp_path / "template.pptx"
    shutil.copyfile(FIXTURE_PATH, dest)
    return str(dest)


def _title(text: str) -> TitleComponent:
    return TitleComponent(text=text)


def _mock_llm(findings: SemanticFindings) -> AsyncMock:
    llm = AsyncMock()
    llm.complete_structured = AsyncMock(return_value=findings)
    return llm


async def test_clean_deck_passes(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)

    slides = [
        SlideIR(
            slide_index=i,
            layout_type=LayoutType.CONTENT_1COL,
            title=_title(f"Slide {i} conclusion"),
            components=[BulletBlock(items=[BulletItem(text="A short bullet point")])],
        )
        for i in range(10)
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)
    build_result = PptxBuilder().build(template_path, manifest, ir)

    report = await run_full_audit(
        build_result=build_result,
        ir=ir,
        manifest=manifest,
        brief="A brief about quarterly performance.",
        llm=_mock_llm(SemanticFindings(findings=[])),
        registry=PromptRegistry("skills"),
        model="qwen",
    )

    assert report.passed is True
    assert report.issues == []


async def test_dirty_deck_flags_multiple_issue_types(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)

    layout = manifest.find_layout(LayoutType.CONTENT_2COL)
    body_idxs = [s.placeholder_idx for s in layout.slots if s.placeholder_type == PlaceholderType.BODY]
    assert len(body_idxs) >= 2

    # the builder now reports each populated placeholder's own real, inherited geometry
    # (Bug 1 fix) rather than trusting the parsed manifest, so a genuine collision must be
    # forced onto the actual layout XML, not just the in-memory manifest, to still trip it
    prs = Presentation(template_path)
    pptx_layout = prs.slide_layouts[layout.layout_index]
    body_placeholders = [p for p in pptx_layout.placeholders if p.placeholder_format.idx in body_idxs]
    anchor = body_placeholders[0]
    for placeholder in body_placeholders[1:]:
        placeholder.left, placeholder.top = anchor.left, anchor.top
        placeholder.width, placeholder.height = anchor.width, anchor.height
    prs.save(template_path)

    manifest = TemplateParser().parse(template_path)

    slides = [
        SlideIR(
            slide_index=0,
            layout_type=LayoutType.CONTENT_2COL,
            title=_title("Overlap slide"),
            components=[
                BulletBlock(items=[BulletItem(text="First block")]),
                BulletBlock(items=[BulletItem(text="Second block")]),
            ],
        ),
        SlideIR(
            slide_index=1,
            layout_type=LayoutType.CONTENT_1COL,
            title=_title("Slide with sneaky junk text"),
            components=[BulletBlock(items=[BulletItem(text="TODO fix this slide")])],
        ),
        *[
            SlideIR(
                slide_index=i,
                layout_type=LayoutType.CONTENT_1COL,
                title=_title(f"Slide {i} conclusion"),
                components=[BulletBlock(items=[BulletItem(text="A short bullet point")])],
            )
            for i in range(2, 10)
        ],
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)
    build_result = PptxBuilder().build(template_path, manifest, ir)

    report = await run_full_audit(
        build_result=build_result,
        ir=ir,
        manifest=manifest,
        brief="A brief about quarterly performance.",
        llm=_mock_llm(SemanticFindings(findings=[])),
        registry=PromptRegistry("skills"),
        model="qwen",
    )

    assert report.passed is False
    issue_types = {i.issue_type for i in report.issues}
    assert IssueType.COLLISION in issue_types
    assert IssueType.PLACEHOLDER_TEXT in issue_types
