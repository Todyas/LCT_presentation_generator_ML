"""Cross-template invariants for app/core/builder/pptx_builder.py.

Builds real .pptx files with python-pptx (this is the expensive part of the
suite), so the template x IR matrix here is deliberately kept small — the
per-component-type render mechanics already have dedicated coverage in
tests/unit/test_shape_factory.py and tests/unit/test_pptx_builder_placeholders.py.
"""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.core.builder.pptx_builder import PptxBuilder
from app.core.parser.template_parser import TemplateParser
from app.models.presentation_ir import MetricCard
from app.models.template_manifest import LayoutType
from tests.factories import make_ir_with_all_component_types, make_slide

TEMPLATE_FIXTURES = [
    "default_template",
    "placeholderless_template",
    "asymmetric_template",
]


def _build(template_path: str):
    manifest = TemplateParser().parse(template_path)
    ir = make_ir_with_all_component_types(template_source_hash=manifest.source_hash)
    result = PptxBuilder().build(template_path, manifest, ir)
    return manifest, ir, result


@pytest.mark.parametrize("template_fixture", TEMPLATE_FIXTURES)
def test_slide_count_matches_ir_and_file_reopens_cleanly(
    tmp_cwd, template_fixture, request
):
    template_path = request.getfixturevalue(template_fixture)
    _, ir, result = _build(template_path)

    reopened = Presentation(result.pptx_path)

    assert len(reopened.slides) == len(ir.slides)


@pytest.mark.parametrize("template_fixture", TEMPLATE_FIXTURES)
def test_output_is_native_only_no_picture_shapes_charts_and_tables_are_native(
    tmp_cwd, template_fixture, request
):
    template_path = request.getfixturevalue(template_fixture)
    _, ir, result = _build(template_path)
    reopened = Presentation(result.pptx_path)

    for slide in reopened.slides:
        for shape in slide.shapes:
            assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE

    chart_slide_indices = {
        s.slide_index for s in ir.slides if any(c.type == "chart" for c in s.components)
    }
    table_slide_indices = {
        s.slide_index for s in ir.slides if any(c.type == "table" for c in s.components)
    }
    for position, slide in enumerate(reopened.slides):
        ir_slide = ir.slides[position]
        if ir_slide.slide_index in chart_slide_indices:
            assert any(
                shape.has_chart for shape in slide.shapes if hasattr(shape, "has_chart")
            )
        if ir_slide.slide_index in table_slide_indices:
            assert any(
                shape.has_table for shape in slide.shapes if hasattr(shape, "has_table")
            )


@pytest.mark.parametrize("template_fixture", TEMPLATE_FIXTURES)
def test_no_ghost_placeholders_remain_empty(tmp_cwd, template_fixture, request):
    template_path = request.getfixturevalue(template_fixture)
    _, _, result = _build(template_path)
    reopened = Presentation(result.pptx_path)

    for slide in reopened.slides:
        for shape in slide.placeholders:
            has_chart = getattr(shape, "has_chart", False)
            has_table = getattr(shape, "has_table", False)
            has_text = shape.has_text_frame and shape.text_frame.text.strip() != ""
            assert has_chart or has_table or has_text, (
                f"empty ghost placeholder idx={shape.placeholder_format.idx} left in output"
            )


@pytest.mark.parametrize("template_fixture", TEMPLATE_FIXTURES)
def test_every_ir_title_appears_on_its_slide(tmp_cwd, template_fixture, request):
    template_path = request.getfixturevalue(template_fixture)
    _, ir, result = _build(template_path)
    reopened = Presentation(result.pptx_path)

    for position, slide in enumerate(reopened.slides):
        all_text = " ".join(
            shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
        )
        assert ir.slides[position].title.text in all_text


def test_placeholderless_template_example_content_does_not_leak_into_output(
    tmp_cwd, placeholderless_template
):
    _, _, result = _build(placeholderless_template)
    reopened = Presentation(result.pptx_path)

    for slide in reopened.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            assert "Example heading" not in shape.text_frame.text
            assert "Example body content" not in shape.text_frame.text


@pytest.mark.parametrize("template_fixture", TEMPLATE_FIXTURES)
def test_every_bbox_lies_within_slide_bounds_and_keys_match_slide_indices(
    tmp_cwd, template_fixture, request
):
    template_path = request.getfixturevalue(template_fixture)
    manifest, ir, result = _build(template_path)

    assert set(result.bbox_map.keys()) == {slide.slide_index for slide in ir.slides}
    for shape_boxes in result.bbox_map.values():
        for box in shape_boxes.values():
            assert box.x >= 0
            assert box.y >= 0
            assert box.x + box.w <= manifest.slide_width_emu
            assert box.y + box.h <= manifest.slide_height_emu


def test_asymmetric_template_components_never_overlap_the_artwork_half(
    tmp_cwd, asymmetric_template
):
    # The synthetic template also carries python-pptx's 11 default full-width
    # layouts alongside the one asymmetric Blank layout, so the IR must
    # explicitly request BLANK to actually land on the constrained layout —
    # CONTENT_1COL would resolve to the unrelated full-width "Title and
    # Content" layout and trivially "pass" this check for the wrong reason.
    from app.models.presentation_ir import PresentationIR

    manifest = TemplateParser().parse(asymmetric_template)
    slides = [make_slide(LayoutType.BLANK, slide_index=0)] + [
        make_slide(LayoutType.BLANK, slide_index=i) for i in range(1, 10)
    ]
    ir = PresentationIR(
        variant="A", template_source_hash=manifest.source_hash, slides=slides
    )
    result = PptxBuilder().build(asymmetric_template, manifest, ir)
    artwork_left_emu = int(manifest.slide_width_emu * 0.5)

    saw_a_component = False
    for slide_index, shape_boxes in result.bbox_map.items():
        for shape_id, box in shape_boxes.items():
            if shape_id == "title":
                continue
            saw_a_component = True
            assert box.x + box.w <= artwork_left_emu, (
                f"slide {slide_index} component {shape_id!r} overlaps the artwork half: "
                f"x={box.x} w={box.w} artwork_left={artwork_left_emu}"
            )
    assert saw_a_component


# --------------------------------------------------------------------------- #
# MetricCard groups: 1..4 cards render pairwise non-overlapping, left-to-right
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize("card_count", [1, 2, 3, 4])
def test_metric_card_group_boxes_are_non_overlapping_and_left_to_right(
    tmp_cwd, default_template, card_count
):
    manifest = TemplateParser().parse(default_template)
    slide = make_slide(
        LayoutType.KPI_DASHBOARD,
        slide_index=0,
        components=[
            MetricCard(label=f"Метрика {i}", value=str(i)) for i in range(card_count)
        ],
    )
    from app.models.presentation_ir import PresentationIR

    filler = [make_slide(LayoutType.CONTENT_1COL, slide_index=i) for i in range(1, 10)]
    ir = PresentationIR(
        variant="A", template_source_hash=manifest.source_hash, slides=[slide, *filler]
    )
    result = PptxBuilder().build(default_template, manifest, ir)

    boxes = [
        box
        for key, box in sorted(result.bbox_map[0].items())
        if key.startswith("component_")
    ]
    assert len(boxes) == card_count
    for i in range(len(boxes) - 1):
        assert (
            boxes[i].x + boxes[i].w <= boxes[i + 1].x
        )  # left-to-right, non-overlapping
