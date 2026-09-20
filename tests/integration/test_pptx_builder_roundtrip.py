import hashlib
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from app.core.builder.pptx_builder import PptxBuilder
from app.core.parser.template_parser import TemplateParser
from app.models.presentation_ir import (
    BulletBlock,
    BulletItem,
    ChartData,
    ChartSeries,
    ImagePlaceholder,
    MetricCard,
    PresentationIR,
    SlideIR,
    TableData,
    TitleComponent,
)
from app.models.template_manifest import (
    Geometry,
    LayoutManifest,
    LayoutSlot,
    LayoutType,
    NormalizedGeometry,
    PlaceholderType,
)

FIXTURE_PATH = "tests/fixtures/templates/generated_minimal.pptx"


def _cache_path_for(pptx_path: str) -> Path:
    file_bytes = Path(pptx_path).read_bytes()
    source_hash = hashlib.sha256(file_bytes).hexdigest()
    return Path(".cache") / f"{source_hash}.manifest.json"


@pytest.fixture(autouse=True)
def _clear_cache():
    cache_path = _cache_path_for(FIXTURE_PATH)
    cache_path.unlink(missing_ok=True)
    yield
    cache_path.unlink(missing_ok=True)


def _copy_template(tmp_path: Path) -> str:
    dest = tmp_path / "template.pptx"
    shutil.copyfile(FIXTURE_PATH, dest)
    return str(dest)


def _title(text: str) -> TitleComponent:
    return TitleComponent(text=text)


def _filler_slides(start_index: int, count: int) -> list[SlideIR]:
    return [
        SlideIR(slide_index=start_index + i, layout_type=LayoutType.SECTION_HEADER, title=_title(f"Filler {i}"))
        for i in range(count)
    ]


def test_build_produces_openable_pptx_with_correct_slide_count(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)

    slides = [
        SlideIR(slide_index=0, layout_type=LayoutType.TITLE_SLIDE, title=_title("Welcome"), components=[]),
        SlideIR(
            slide_index=1,
            layout_type=LayoutType.SECTION_HEADER,
            title=_title("Agenda"),
            components=[BulletBlock(items=[BulletItem(text="Intro"), BulletItem(text="Results")])],
        ),
        SlideIR(
            slide_index=2,
            layout_type=LayoutType.CONTENT_2COL,
            title=_title("Two Streams"),
            components=[
                BulletBlock(items=[BulletItem(text="Stream A")]),
                BulletBlock(items=[BulletItem(text="Stream B")]),
            ],
        ),
        SlideIR(
            slide_index=3,
            layout_type=LayoutType.COMPARISON,
            title=_title("Before and After"),
            components=[
                BulletBlock(items=[BulletItem(text="Old process")]),
                BulletBlock(items=[BulletItem(text="New process")]),
                TableData(headers=["Metric", "Value"], rows=[["Speed", "Fast"]]),
            ],
        ),
        SlideIR(slide_index=4, layout_type=LayoutType.BLANK, title=_title("Divider"), components=[]),
        SlideIR(
            slide_index=5,
            layout_type=LayoutType.CONTENT_1COL,
            title=_title("Product Shot"),
            components=[ImagePlaceholder(alt_text="hero image", role="hero")],
        ),
        SlideIR(
            slide_index=6,
            layout_type=LayoutType.TITLE_SLIDE,
            title=_title("Growth"),
            components=[MetricCard(label="Revenue", value="$1.2M")],
        ),
        SlideIR(
            slide_index=7,
            layout_type=LayoutType.SECTION_HEADER,
            title=_title("Quarterly Trend"),
            components=[
                ChartData(
                    chart_type="column",
                    categories=["Q1", "Q2", "Q3"],
                    series=[ChartSeries(name="Revenue", values=[1.0, 2.0, 3.0])],
                )
            ],
        ),
        SlideIR(
            slide_index=8,
            layout_type=LayoutType.CONTENT_2COL,
            title=_title("Breakdown"),
            components=[TableData(headers=["Region", "Sales"], rows=[["EU", "100"], ["US", "200"]])],
        ),
        SlideIR(
            slide_index=9,
            layout_type=LayoutType.COMPARISON,
            title=_title("Key Metrics"),
            components=[
                MetricCard(label="Churn", value="2%"),
                MetricCard(label="NPS", value="45"),
            ],
        ),
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)

    result = PptxBuilder().build(template_path, manifest, ir)

    reopened = Presentation(result.pptx_path)
    assert len(reopened.slides._sldIdLst) == 10
    assert len(list(reopened.slides)) == 10


def test_image_placeholder_renders_even_without_picture_slot(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)
    layout = manifest.find_layout_or_fallback(LayoutType.TITLE_SLIDE)
    assert layout.slot_by_type(PlaceholderType.PICTURE) is None

    slides = [
        SlideIR(
            slide_index=0,
            layout_type=LayoutType.TITLE_SLIDE,
            title=_title("No Picture Slot Here"),
            components=[ImagePlaceholder(alt_text="a placeholder image", role="illustration")],
        ),
        *_filler_slides(1, 9),
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)

    result = PptxBuilder().build(template_path, manifest, ir)

    reopened = Presentation(result.pptx_path)
    slide = reopened.slides[0]
    assert len(list(slide.shapes)) >= 2


def test_duplicate_component_types_do_not_collide(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)
    layout = manifest.find_layout_or_fallback(LayoutType.CONTENT_1COL)
    body_slots = [s for s in layout.slots if s.placeholder_type == PlaceholderType.BODY]
    assert len(body_slots) == 1

    slides = [
        SlideIR(
            slide_index=0,
            layout_type=LayoutType.CONTENT_1COL,
            title=_title("Duplicate Bullets"),
            components=[
                BulletBlock(items=[BulletItem(text="First block")]),
                BulletBlock(items=[BulletItem(text="Second block")]),
            ],
        ),
        *_filler_slides(1, 9),
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)

    result = PptxBuilder().build(template_path, manifest, ir)

    # exactly two BulletBlocks sharing one slot are split into non-overlapping side-by-side
    # columns rather than either colliding at the same geometry or stacking vertically
    boxes = result.bbox_map[0]
    first_bbox = boxes["component_0"]
    second_bbox = boxes["component_1"]
    assert first_bbox.y == second_bbox.y
    assert first_bbox.x != second_bbox.x
    assert first_bbox.x + first_bbox.w <= second_bbox.x


def test_reopened_file_has_no_repair_warning_indicators(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)

    slides = [
        SlideIR(
            slide_index=i,
            layout_type=LayoutType.SECTION_HEADER,
            title=_title(f"Slide {i}"),
            components=[BulletBlock(items=[BulletItem(text="Item one")])],
        )
        for i in range(10)
    ]
    ir = PresentationIR(variant="B", template_source_hash=manifest.source_hash, slides=slides)

    result = PptxBuilder().build(template_path, manifest, ir)

    Presentation(result.pptx_path)


def test_fallback_geometry_stays_within_asymmetric_layout_content_region(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)
    slide_width = manifest.slide_width_emu
    slide_height = manifest.slide_height_emu

    right_half_geometry = Geometry(
        left_emu=int(slide_width * 0.55), top_emu=int(slide_height * 0.2),
        width_emu=int(slide_width * 0.4), height_emu=int(slide_height * 0.6),
    )
    right_half_layout = LayoutManifest(
        layout_index=1,
        layout_name="Right Half Custom",
        layout_type=LayoutType.KPI_DASHBOARD,
        slots=[
            LayoutSlot(
                placeholder_idx=1,
                placeholder_type=PlaceholderType.BODY,
                geometry=right_half_geometry,
                normalized=NormalizedGeometry(x=0.55, y=0.2, w=0.4, h=0.6),
                name="Body Right",
            ),
        ],
    )
    asymmetric_manifest = manifest.model_copy(update={"layouts": [right_half_layout]})

    slides = [
        SlideIR(
            slide_index=0,
            layout_type=LayoutType.KPI_DASHBOARD,
            title=_title("Asymmetric Template"),
            components=[TableData(headers=["A", "B"], rows=[["1", "2"]])],
        ),
        *_filler_slides(1, 9),
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)

    result = PptxBuilder().build(template_path, asymmetric_manifest, ir)

    bbox = result.bbox_map[0]["component_0"]
    assert bbox.x >= 0.5 * slide_width
    assert bbox.x + bbox.w <= slide_width
