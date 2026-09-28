from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from app.core.builder.shape_factory import (
    render_bullet_block,
    render_bullet_component,
    render_chart,
    render_comparison,
    render_icon_list,
    render_metric_card,
    render_metric_card_group,
    render_process,
    render_table,
)
from app.core.parser.font_resolver import resolve_font_path
from app.models.presentation_ir import (
    BulletBlock,
    BulletItem,
    ChartData,
    ChartSeries,
    ComparisonData,
    IconListData,
    IconListItem,
    MetricCard,
    ProcessData,
    ProcessStep,
    TableData,
)
from app.models.template_manifest import Geometry, ThemeColors

THEME = ThemeColors(
    dk1="000000",
    lt1="FFFFFF",
    dk2="1F1F1F",
    lt2="E7E7E7",
    accent1="4472C4",
    accent2="ED7D31",
    accent3="A5A5A5",
    accent4="FFC000",
    accent5="5B9BD5",
    accent6="70AD47",
    hlink="0563C1",
    fol_hlink="954F72",
)
FONT_PATH = resolve_font_path(None)


def _new_slide():
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[0])


def test_render_chart_maps_pie_type():
    slide = _new_slide()
    geometry = Geometry(
        left_emu=0, top_emu=0, width_emu=4_000_000, height_emu=3_000_000
    )
    chart = ChartData(
        chart_type="pie",
        categories=["A", "B"],
        series=[ChartSeries(name="s1", values=[1.0, 2.0])],
    )

    render_chart(slide, geometry, chart)

    assert slide.shapes[-1].chart.chart_type == XL_CHART_TYPE.PIE


def test_render_chart_uses_template_accent_palette():
    slide = _new_slide()
    geometry = Geometry(
        left_emu=0, top_emu=0, width_emu=4_000_000, height_emu=3_000_000
    )
    chart = ChartData(
        chart_type="column",
        categories=["A", "B"],
        series=[ChartSeries(name="s1", values=[1.0, 2.0])],
    )

    render_chart(slide, geometry, chart, THEME)

    assert (
        str(slide.shapes[-1].chart.series[0].format.fill.fore_color.rgb)
        == THEME.accent1
    )


def test_render_table_fills_headers_and_rows():
    slide = _new_slide()
    geometry = Geometry(
        left_emu=0, top_emu=0, width_emu=4_000_000, height_emu=2_000_000
    )
    table = TableData(headers=["H1", "H2"], rows=[["r1c1", "r1c2"], ["r2c1", "r2c2"]])

    render_table(slide, geometry, table)

    pptx_table = slide.shapes[-1].table
    assert pptx_table.cell(0, 0).text == "H1"
    assert pptx_table.cell(1, 0).text == "r1c1"


def test_render_metric_card_returns_requested_bbox():
    slide = _new_slide()
    geometry = Geometry(
        left_emu=100_000, top_emu=200_000, width_emu=1_500_000, height_emu=900_000
    )
    card = MetricCard(label="Revenue", value="$1.2M")

    bbox = render_metric_card(slide, geometry, card, THEME, FONT_PATH)

    assert (bbox.x, bbox.y, bbox.w, bbox.h) == (
        geometry.left_emu,
        geometry.top_emu,
        geometry.width_emu,
        geometry.height_emu,
    )


def test_render_metric_card_autofits_long_value_text():
    slide = _new_slide()
    geometry = Geometry(left_emu=0, top_emu=0, width_emu=600_000, height_emu=900_000)
    card = MetricCard(label="Total", value="123 456 789")

    render_metric_card(slide, geometry, card, THEME, FONT_PATH)

    value_box = slide.shapes[-2]
    run_size = value_box.text_frame.paragraphs[0].runs[0].font.size
    assert run_size < Pt(32)


def test_render_metric_card_group_produces_n_non_overlapping_boxes():
    for n in (2, 3):
        slide = _new_slide()
        geometry = Geometry(
            left_emu=0, top_emu=0, width_emu=6_000_000, height_emu=900_000
        )
        cards = [MetricCard(label=f"Metric {i}", value=str(i)) for i in range(n)]

        boxes = render_metric_card_group(slide, geometry, cards, THEME, FONT_PATH)

        assert len(boxes) == n
        for i in range(n - 1):
            assert boxes[i].x + boxes[i].w <= boxes[i + 1].x
            assert boxes[i].x < boxes[i + 1].x


def test_render_bullet_block_creates_one_paragraph_per_item():
    slide = _new_slide()
    textbox = slide.shapes.add_textbox(0, 0, 4_000_000, 2_000_000)
    block = BulletBlock(
        items=[
            BulletItem(text="first"),
            BulletItem(text="second"),
            BulletItem(text="third"),
        ]
    )

    render_bullet_block(
        textbox.text_frame, block, font_path="unused", font_name="Calibri"
    )

    assert len(textbox.text_frame.paragraphs) == 3
    assert [p.text for p in textbox.text_frame.paragraphs] == [
        "first",
        "second",
        "third",
    ]


def test_fallback_bullet_component_is_transparent_text_not_a_panel():
    slide = _new_slide()
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
    geometry = Geometry(
        left_emu=100_000,
        top_emu=200_000,
        width_emu=4_000_000,
        height_emu=2_000_000,
    )
    block = BulletBlock(
        items=[BulletItem(text="Main conclusion"), BulletItem(text="Supporting detail")]
    )

    render_bullet_component(
        slide,
        geometry,
        None,
        block,
        FONT_PATH,
        "Calibri",
        THEME,
        9_144_000,
        6_858_000,
    )

    shape = slide.shapes[-1]
    assert shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
    assert shape.fill.type != MSO_FILL_TYPE.SOLID
    assert str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert shape.text_frame.paragraphs[0].runs[0].font.bold is True


def test_semantic_visuals_render_as_native_powerpoint_shapes():
    slide = _new_slide()
    geometry = Geometry(
        left_emu=100_000, top_emu=300_000, width_emu=7_000_000, height_emu=3_000_000
    )

    render_comparison(
        slide,
        geometry,
        ComparisonData(
            left_title="До",
            left_items=["Ручной процесс"],
            right_title="После",
            right_items=["Автоматизация"],
        ),
        THEME,
        "Arial",
        9_144_000,
        6_858_000,
    )
    comparison_shape_count = len(slide.shapes)
    render_process(
        slide,
        geometry,
        ProcessData(
            steps=[ProcessStep(title=value) for value in ("Сбор", "Анализ", "Экспорт")]
        ),
        THEME,
        "Arial",
        9_144_000,
        6_858_000,
    )
    render_icon_list(
        slide,
        geometry,
        IconListData(
            items=[IconListItem(title="Скорость"), IconListItem(title="Контроль")]
        ),
        THEME,
        "Arial",
        9_144_000,
        6_858_000,
    )

    assert comparison_shape_count >= 6
    assert len(slide.shapes) > comparison_shape_count
