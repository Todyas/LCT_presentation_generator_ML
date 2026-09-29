import copy
import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu

from app.core.builder.pptx_builder import PptxBuilder
from app.core.parser.template_parser import TemplateParser
from app.models.presentation_ir import (
    PresentationIR,
    ProcessData,
    ProcessStep,
    SlideIR,
    TableData,
    TitleComponent,
)
from app.models.template_manifest import LayoutType

SLIDE_W = 12_192_000
SLIDE_H = 6_858_000
PICTURE_LEFT = int(SLIDE_W * 0.4)


def _make_template(path) -> str:
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    layout = prs.slide_layouts.get_by_name("Title Only")
    layout.name = "Art Right"

    for placeholder in layout.placeholders:
        if placeholder.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        ):
            placeholder.left = Emu(int(SLIDE_W * 0.05))
            placeholder.top = Emu(int(SLIDE_H * 0.04))
            placeholder.width = Emu(int(SLIDE_W * 0.35))
            placeholder.height = Emu(int(SLIDE_H * 0.11))

    # a decorative picture over the right 60% of the layout
    buffer = io.BytesIO()
    Image.new("RGB", (40, 40), (30, 60, 200)).save(buffer, format="PNG")
    buffer.seek(0)
    scratch = prs.slides.add_slide(prs.slide_layouts.get_by_name("Blank"))
    picture = scratch.shapes.add_picture(
        buffer,
        Emu(PICTURE_LEFT),
        Emu(0),
        Emu(SLIDE_W - PICTURE_LEFT),
        Emu(SLIDE_H),
    )
    _, rid = layout.part.get_or_add_image_part(io.BytesIO(buffer.getvalue()))
    element = copy.deepcopy(picture._element)
    blip = element.xpath(".//a:blip")[0]
    blip.set(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed",
        rid,
    )
    layout.shapes._spTree.append(element)
    slide_id = prs.slides._sldIdLst[0]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)

    out = str(path / "art_right.pptx")
    prs.save(out)
    return out


@pytest.fixture()
def parsed(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    template = _make_template(tmp_path)
    manifest = TemplateParser().parse(template)
    layout = next(l for l in manifest.layouts if l.layout_name == "Art Right")
    return template, manifest, layout


def test_content_region_avoids_picture(parsed):
    _, _, layout = parsed
    region = layout.content_region
    assert region is not None
    assert layout.decoration_coverage > 0.5
    assert region.left_emu >= int(SLIDE_W * 0.05) - 1
    assert region.right_emu <= PICTURE_LEFT
    assert region.right_emu <= int(SLIDE_W * 0.4)


def test_builder_places_components_inside_region_below_title(parsed):
    template, manifest, layout = parsed
    region = layout.content_region
    assert region is not None

    slides = [
        SlideIR(
            slide_index=0,
            layout_type=LayoutType.CONTENT_1COL,
            layout_index=layout.layout_index,
            title=TitleComponent(text="Сравнение вариантов запуска"),
            components=[
                TableData(
                    headers=["Этап", "Срок"],
                    rows=[["Подготовка", "1 нед"], ["Запуск", "2 нед"]],
                ),
                ProcessData(
                    steps=[
                        ProcessStep(title="Идея", description="Сбор требований"),
                        ProcessStep(title="Пилот", description="Проверка гипотез"),
                        ProcessStep(title="Запуск", description="Выход на рынок"),
                    ]
                ),
            ],
        ),
        *[
            SlideIR(
                slide_index=i,
                layout_type=LayoutType.SECTION_HEADER,
                title=TitleComponent(text=f"Раздел {i}"),
            )
            for i in range(1, 10)
        ],
    ]
    ir = PresentationIR(
        variant="A", template_source_hash=manifest.source_hash, slides=slides
    )
    result = PptxBuilder().build(template, manifest, ir)
    boxes = result.bbox_map[0]

    title = boxes["title"]
    components = [boxes["component_0"], boxes["component_1"]]
    for box in components:
        assert box.x >= region.left_emu
        assert box.x + box.w <= region.right_emu
        assert box.y >= title.y + title.h
        assert box.y + box.h <= region.bottom_emu + 1
        assert box.x + box.w <= PICTURE_LEFT
    first, second = components
    assert first.y + first.h <= second.y
