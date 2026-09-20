from pptx import Presentation
from pptx.util import Inches

from app.core.auditor.density_audit import (
    run_density_audit_on_ir,
    run_placeholder_text_audit,
)
from app.models.presentation_ir import (
    BulletBlock,
    BulletItem,
    PresentationIR,
    SlideIR,
    TitleComponent,
)
from app.models.template_manifest import LayoutType


def _build_pptx_with_textbox(tmp_path, text: str) -> str:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = text
    path = str(tmp_path / "deck.pptx")
    prs.save(path)
    return path


def test_placeholder_regex_ignores_false_positive_substring(tmp_path):
    path = _build_pptx_with_textbox(tmp_path, "Maximum efficiency this quarter")

    issues = run_placeholder_text_audit(path)

    assert issues == []


def test_placeholder_regex_detects_todo(tmp_path):
    path = _build_pptx_with_textbox(tmp_path, "TODO: fix this slide")

    issues = run_placeholder_text_audit(path)

    assert len(issues) == 1
    assert issues[0].issue_type.value == "PLACEHOLDER_TEXT"


def _valid_ir() -> PresentationIR:
    slides = [
        SlideIR(
            slide_index=i,
            layout_type=LayoutType.CONTENT_1COL,
            title=TitleComponent(text=f"Slide {i} conclusion"),
            components=[BulletBlock(items=[BulletItem(text="A short bullet point")])],
        )
        for i in range(10)
    ]
    return PresentationIR(variant="A", template_source_hash="abc", slides=slides)


def test_density_audit_on_ir_clean_input_returns_empty():
    ir = _valid_ir()

    issues = run_density_audit_on_ir(ir)

    assert issues == []
