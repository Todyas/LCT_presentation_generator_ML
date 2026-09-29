import io
import zipfile

from lxml import etree
from PIL import Image
from pptx import Presentation

from app.core.builder.shape_factory import (
    _background_hex,
    _contrasting_text_hex,
    _is_dark,
    _surface_palette,
)
from app.core.parser.template_parser import TemplateParser
from app.models.template_manifest import ThemeColors

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

THEME = ThemeColors(
    dk1="000000",
    lt1="FFFFFF",
    dk2="1C1D22",
    lt2="E7E7E7",
    accent1="FF0053",
    accent2="ED7D31",
    accent3="A5A5A5",
    accent4="FFC000",
    accent5="5B9BD5",
    accent6="70AD47",
    hlink="0563C1",
    fol_hlink="954F72",
)


def _png(color: tuple[int, int, int]) -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (100, 100), color).save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _set_master_background_picture(prs, color: tuple[int, int, int]) -> None:
    master = prs.slide_master
    _, rel_id = master.part.get_or_add_image_part(_png(color))
    background = etree.fromstring(
        f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:bgPr>'
        f'<a:blipFill><a:blip r:embed="{rel_id}"/>'
        "<a:stretch><a:fillRect/></a:stretch></a:blipFill>"
        "<a:effectLst/></p:bgPr></p:bg>"
    )
    c_sld = master._element.find(f"{{{NS_P}}}cSld")
    old = c_sld.find(f"{{{NS_P}}}bg")
    if old is not None:
        c_sld.remove(old)
    c_sld.insert(0, background)


def test_master_theme_wins_over_unused_theme1(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "base.pptx"
    Presentation().save(source)

    patched = tmp_path / "themed.pptx"
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(patched, "w") as zout:
        theme1 = zin.read("ppt/theme/theme1.xml").decode("utf-8")
        assert "4F81BD" in theme1 and "1F497D" in theme1
        theme2 = theme1.replace("4F81BD", "FF0053").replace("1F497D", "1C1D22")
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/slideMasters/_rels/slideMaster1.xml.rels":
                data = data.replace(b"theme/theme1.xml", b"theme/theme2.xml")
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/ppt/theme/theme2.xml" ContentType="'
                    b'application/vnd.openxmlformats-officedocument.theme+xml"/>'
                    b"</Types>",
                )
            zout.writestr(item, data)
        zout.writestr("ppt/theme/theme2.xml", theme2)

    manifest = TemplateParser().parse(str(patched))

    assert manifest.colors.accent1 == "FF0053"
    assert manifest.colors.dk2 == "1C1D22"


def test_dark_picture_background_gives_light_text():
    prs = Presentation()
    _set_master_background_picture(prs, (85, 25, 117))
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    background = _background_hex(slide, THEME, prs.slide_width, prs.slide_height)

    assert _is_dark(background)
    assert _contrasting_text_hex(background, THEME) == THEME.lt1
    assert slide._element.find(f"{{{NS_P}}}cSld/{{{NS_P}}}bg") is None


def test_dark_background_surface_is_not_white():
    prs = Presentation()
    _set_master_background_picture(prs, (85, 25, 117))
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    surface, text, accent = _surface_palette(
        slide, THEME, prs.slide_width, prs.slide_height
    )

    assert _is_dark(surface)
    assert text == THEME.lt1
    assert accent == THEME.accent1


def test_light_default_background_keeps_dark_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    background = _background_hex(slide, THEME, prs.slide_width, prs.slide_height)
    surface, text, _ = _surface_palette(slide, THEME, prs.slide_width, prs.slide_height)

    assert not _is_dark(background)
    assert not _is_dark(surface)
    assert text == THEME.dk1


def test_layout_picture_background_overrides_master_solid_fill():
    prs = Presentation()
    layout = prs.slide_layouts[6]
    _, rel_id = layout.part.get_or_add_image_part(_png((10, 10, 40)))
    background = etree.fromstring(
        f'<p:bg xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><p:bgPr>'
        f'<a:blipFill><a:blip r:embed="{rel_id}"/></a:blipFill>'
        "<a:effectLst/></p:bgPr></p:bg>"
    )
    layout._element.find(f"{{{NS_P}}}cSld").insert(0, background)
    slide = prs.slides.add_slide(layout)

    assert _is_dark(_background_hex(slide, THEME, prs.slide_width, prs.slide_height))
