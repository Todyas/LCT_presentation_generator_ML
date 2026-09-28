from pptx import Presentation
from pptx.util import Inches

from app.core.parser.template_parser import TemplateParser
from app.models.template_manifest import PlaceholderType


def test_parser_infers_content_slots_from_example_slide_text_boxes(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(0.8), Inches(0.4), Inches(8), Inches(0.7)
    ).text_frame.text = "Содержательный заголовок"
    slide.shapes.add_textbox(
        Inches(0.8), Inches(1.6), Inches(5.4), Inches(4)
    ).text_frame.text = "Основной контент"
    path = tmp_path / "shape-template.pptx"
    prs.save(path)

    manifest = TemplateParser().parse(str(path))
    blank_layout = next(
        layout for layout in manifest.layouts if layout.layout_index == 6
    )

    assert any(
        slot.inferred and slot.placeholder_type == PlaceholderType.TITLE
        for slot in blank_layout.slots
    )
    assert any(
        slot.inferred and slot.placeholder_type == PlaceholderType.BODY
        for slot in blank_layout.slots
    )
