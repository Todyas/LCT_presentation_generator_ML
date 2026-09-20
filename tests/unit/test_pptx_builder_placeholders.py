import hashlib
import shutil
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from app.core.builder.pptx_builder import PptxBuilder
from app.core.parser.template_parser import TemplateParser
from app.models.presentation_ir import (
    BulletBlock,
    BulletItem,
    PresentationIR,
    SlideIR,
    TableData,
    TitleComponent,
)
from app.models.template_manifest import LayoutType, PlaceholderType

FIXTURE_PATH = "tests/fixtures/templates/generated_minimal.pptx"


def _cache_path_for(pptx_path: str) -> Path:
    file_bytes = Path(pptx_path).read_bytes()
    source_hash = hashlib.sha256(file_bytes).hexdigest()
    return Path(".cache") / f"{source_hash}.manifest.json"


@pytest.fixture(autouse=True)
def _clear_cache():
    cache_path = _cache_path_for(FIXTURE_PATH)
    cache_path.unlink(missing_ok=True)
    yield
    cache_path.unlink(missing_ok=True)


def _copy_template(tmp_path: Path) -> str:
    dest = tmp_path / "template.pptx"
    shutil.copyfile(FIXTURE_PATH, dest)
    return str(dest)


def _mutate_layout_placeholder_type(template_path: str, layout_index: int, idx: int, new_type: PP_PLACEHOLDER) -> None:
    prs = Presentation(template_path)
    layout = prs.slide_layouts[layout_index]
    placeholder = next(p for p in layout.placeholders if p.placeholder_format.idx == idx)
    placeholder.element.ph.type = new_type
    prs.save(template_path)


def _title(text: str) -> TitleComponent:
    return TitleComponent(text=text)


def _filler_slides(start_index: int, count: int) -> list[SlideIR]:
    return [
        SlideIR(slide_index=start_index + i, layout_type=LayoutType.SECTION_HEADER, title=_title(f"Filler {i}"))
        for i in range(count)
    ]


def test_unused_placeholder_is_removed(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)
    layout = manifest.find_layout_or_fallback(LayoutType.TITLE_SLIDE)
    subtitle_slot = layout.slot_by_type(PlaceholderType.SUBTITLE)
    assert subtitle_slot is not None

    slides = [
        SlideIR(slide_index=0, layout_type=LayoutType.TITLE_SLIDE, title=_title("No Subtitle Content"), components=[]),
        *_filler_slides(1, 9),
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)

    result = PptxBuilder().build(template_path, manifest, ir)

    reopened = Presentation(result.pptx_path)
    slide = reopened.slides[0]
    remaining_idxs = {s.placeholder_format.idx for s in slide.placeholders}
    assert subtitle_slot.placeholder_idx not in remaining_idxs


def test_bullet_block_populates_existing_placeholder_not_a_new_textbox(tmp_path):
    template_path = _copy_template(tmp_path)
    manifest = TemplateParser().parse(template_path)
    layout = manifest.find_layout_or_fallback(LayoutType.CONTENT_1COL)
    body_slots = [s for s in layout.slots if s.placeholder_type == PlaceholderType.BODY]
    assert len(body_slots) == 1

    slides = [
        SlideIR(
            slide_index=0,
            layout_type=LayoutType.CONTENT_1COL,
            title=_title("Single Body Slot"),
            components=[BulletBlock(items=[BulletItem(text="Only bullet")])],
        ),
        *_filler_slides(1, 9),
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)

    result = PptxBuilder().build(template_path, manifest, ir)

    reopened = Presentation(result.pptx_path)
    slide = reopened.slides[0]
    assert len(list(slide.shapes)) == 2


def test_table_component_uses_insert_table_on_matching_placeholder(tmp_path):
    template_path = _copy_template(tmp_path)
    _mutate_layout_placeholder_type(template_path, layout_index=8, idx=2, new_type=PP_PLACEHOLDER.TABLE)
    manifest = TemplateParser().parse(template_path)
    # mutating a BODY slot to TABLE also reclassifies the layout itself (layout_classifier.py
    # promotes any layout containing a TABLE placeholder to TABLE_FOCUSED)
    layout = manifest.find_layout_or_fallback(LayoutType.TABLE_FOCUSED)
    table_slot = layout.slot_by_type(PlaceholderType.TABLE)
    assert table_slot is not None

    slides = [
        SlideIR(
            slide_index=0,
            layout_type=LayoutType.TABLE_FOCUSED,
            title=_title("Table Slot"),
            components=[TableData(headers=["A", "B"], rows=[["1", "2"], ["3", "4"]])],
        ),
        *_filler_slides(1, 9),
    ]
    ir = PresentationIR(variant="A", template_source_hash=manifest.source_hash, slides=slides)

    result = PptxBuilder().build(template_path, manifest, ir)

    reopened = Presentation(result.pptx_path)
    slide = reopened.slides[0]
    table_shape = next(s for s in slide.shapes if s.has_table)
    # insert_table keeps the placeholder's left/top/width but derives height from row count,
    # so height is intentionally not asserted here (documented python-pptx behavior).
    assert table_shape.left == table_slot.geometry.left_emu
    assert table_shape.top == table_slot.geometry.top_emu
    assert table_shape.width == table_slot.geometry.width_emu
    assert len(list(slide.shapes)) == 2
