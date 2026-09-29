from __future__ import annotations

import hashlib
import logging
import zipfile
from collections import Counter
from pathlib import Path
from statistics import median

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from app.core.parser.layout_classifier import classify_layout
from app.core.parser.theme_extractor import extract_font_scheme, extract_theme_colors
from app.models.template_manifest import (
    BrandProfile,
    FontScheme,
    Geometry,
    LayoutManifest,
    LayoutSlot,
    NormalizedGeometry,
    PlaceholderType,
    TemplateManifest,
    ThemeColors,
)

logger = logging.getLogger(__name__)

_THEME_PATH = "ppt/theme/theme1.xml"

_PLACEHOLDER_TYPE_MAP: dict[int, PlaceholderType] = {
    PP_PLACEHOLDER.TITLE: PlaceholderType.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE: PlaceholderType.TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE: PlaceholderType.TITLE,
    PP_PLACEHOLDER.SUBTITLE: PlaceholderType.SUBTITLE,
    PP_PLACEHOLDER.BODY: PlaceholderType.BODY,
    PP_PLACEHOLDER.VERTICAL_BODY: PlaceholderType.BODY,
    PP_PLACEHOLDER.OBJECT: PlaceholderType.BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT: PlaceholderType.BODY,
    PP_PLACEHOLDER.PICTURE: PlaceholderType.PICTURE,
    PP_PLACEHOLDER.TABLE: PlaceholderType.TABLE,
    PP_PLACEHOLDER.CHART: PlaceholderType.CHART,
    PP_PLACEHOLDER.FOOTER: PlaceholderType.FOOTER,
    PP_PLACEHOLDER.SLIDE_NUMBER: PlaceholderType.SLIDE_NUMBER,
    PP_PLACEHOLDER.DATE: PlaceholderType.DATE,
}


class TemplateParseError(Exception):
    pass


def _direct_rgb(color) -> str | None:
    try:
        return str(color.rgb) if color.rgb is not None else None
    except (AttributeError, ValueError):
        return None


def _extract_brand_profile(prs: Presentation) -> BrandProfile:
    title_sizes: list[float] = []
    body_sizes: list[float] = []
    colors: Counter[str] = Counter()
    owners = [
        *prs.slide_masters,
        *(layout for master in prs.slide_masters for layout in master.slide_layouts),
        *prs.slides,
    ]
    for owner in owners:
        for shape in owner.shapes:
            if hasattr(shape, "fill"):
                try:
                    color = _direct_rgb(shape.fill.fore_color)
                except (AttributeError, TypeError, ValueError):
                    color = None
                if color:
                    colors[color] += 1
            if not getattr(shape, "has_text_frame", False):
                continue
            is_title_region = shape.top < prs.slide_height * 0.3
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is None:
                        continue
                    target = title_sizes if is_title_region else body_sizes
                    target.append(run.font.size.pt)
                    text_color = _direct_rgb(run.font.color)
                    if text_color:
                        colors[text_color] += 1
    title_candidates = [size for size in title_sizes if size >= 18]
    title_size = max(18, min(54, median(title_candidates))) if title_candidates else 32
    body_size = max(10, min(32, median(body_sizes))) if body_sizes else 20
    return BrandProfile(
        title_size_pt=title_size,
        body_size_pt=body_size,
        sampled_colors=[color for color, _ in colors.most_common(12)],
    )


def _inferred_slots(
    owner,
    slide_width: int,
    slide_height: int,
    id_offset: int = 0,
    include_text_content: bool = False,
) -> list[LayoutSlot]:
    """Infer editable content regions from ordinary template shapes.

    Many corporate decks use styled text boxes instead of PowerPoint
    placeholders.  Treat those boxes as semantic slots while keeping their
    geometry separate from the decorative master artwork.
    """

    result: list[LayoutSlot] = []
    for shape in owner.shapes:
        if getattr(shape, "is_placeholder", False):
            continue
        has_chart = bool(getattr(shape, "has_chart", False))
        has_table = bool(getattr(shape, "has_table", False))
        has_text = bool(getattr(shape, "has_text_frame", False))
        if not (has_chart or has_table or has_text):
            continue
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        if width <= 0 or height <= 0:
            continue
        normalized = NormalizedGeometry(
            x=max(0.0, min(1.0, left / slide_width)),
            y=max(0.0, min(1.0, top / slide_height)),
            w=max(0.001, min(1.0, width / slide_width)),
            h=max(0.001, min(1.0, height / slide_height)),
        )
        area = normalized.w * normalized.h
        if area < 0.008 or area > 0.82 or normalized.y > 0.9:
            continue
        text = shape.text_frame.text.strip().casefold() if has_text else ""
        if text and not include_text_content and not (has_chart or has_table):
            continue
        name = (getattr(shape, "name", "") or "").casefold()
        title_hint = any(
            token in f"{name} {text}" for token in ("title", "заголов", "header")
        )
        is_title = title_hint or (normalized.y < 0.25 and normalized.h < 0.28)
        slot_type = (
            PlaceholderType.CHART
            if has_chart
            else PlaceholderType.TABLE
            if has_table
            else PlaceholderType.TITLE
            if is_title
            else PlaceholderType.BODY
        )
        result.append(
            LayoutSlot(
                placeholder_idx=-10_000 - id_offset - int(shape.shape_id),
                placeholder_type=slot_type,
                geometry=Geometry(
                    left_emu=left,
                    top_emu=top,
                    width_emu=width,
                    height_emu=height,
                ),
                normalized=normalized,
                name=getattr(shape, "name", "") or "inferred content region",
                inferred=True,
            )
        )
    return result


def _resolved_placeholder_box(placeholder, master) -> tuple[int, int, int, int] | None:
    values = (placeholder.left, placeholder.top, placeholder.width, placeholder.height)
    if None not in values:
        return values
    idx = placeholder.placeholder_format.idx
    inherited = next(
        (
            candidate
            for candidate in master.placeholders
            if candidate.placeholder_format.idx == idx
        ),
        None,
    )
    if inherited is None:
        return None
    inherited_values = (
        inherited.left,
        inherited.top,
        inherited.width,
        inherited.height,
    )
    return inherited_values if None not in inherited_values else None


def _map_placeholder_type(ph_type: int | None) -> PlaceholderType:
    if ph_type is None:
        return PlaceholderType.OTHER
    return _PLACEHOLDER_TYPE_MAP.get(ph_type, PlaceholderType.OTHER)


def _read_theme_bytes(pptx_path: str) -> bytes:
    with zipfile.ZipFile(pptx_path) as zf:
        if _THEME_PATH not in zf.namelist():
            return b""
        return zf.read(_THEME_PATH)


class TemplateParser:
    def parse(self, pptx_path: str) -> TemplateManifest:
        file_bytes = Path(pptx_path).read_bytes()
        source_hash = hashlib.sha256(file_bytes).hexdigest()

        cache_dir = Path(".cache")
        cache_path = cache_dir / f"{source_hash}.v4.manifest.json"
        cache_writable = True
        try:
            cache_dir.mkdir(exist_ok=True)
        except OSError:
            logger.warning(
                "cache directory %s is not writable, caching disabled", cache_dir
            )
            cache_writable = False

        if cache_writable and cache_path.exists():
            try:
                return TemplateManifest.model_validate_json(cache_path.read_text())
            except OSError:
                logger.warning("cannot read template cache %s; reparsing", cache_path)

        try:
            prs = Presentation(pptx_path)
        except Exception as exc:
            raise TemplateParseError(f"cannot open pptx: {exc}") from exc

        theme_bytes = _read_theme_bytes(pptx_path)
        colors = ThemeColors(**extract_theme_colors(theme_bytes))
        major_latin, minor_latin = extract_font_scheme(theme_bytes)
        fonts = FontScheme(major_latin=major_latin, minor_latin=minor_latin)
        brand_profile = _extract_brand_profile(prs)

        layouts: list[LayoutManifest] = []
        layout_index_by_partname: dict[str, int] = {}
        layout_index = 0
        # prs.slide_layouts only exposes the first slide master's layouts; a corporate
        # template's real content layouts often live on additional masters, so every
        # master must be walked or the manifest silently only sees a handful of title slides
        for master in prs.slide_masters:
            for layout in master.slide_layouts:
                slots: list[LayoutSlot] = []
                for placeholder in layout.placeholders:
                    resolved_box = _resolved_placeholder_box(placeholder, master)
                    if resolved_box is None:
                        continue
                    left, top, width, height = resolved_box
                    geometry = Geometry(
                        left_emu=left,
                        top_emu=top,
                        width_emu=width,
                        height_emu=height,
                    )
                    normalized = NormalizedGeometry(
                        x=left / prs.slide_width,
                        y=top / prs.slide_height,
                        w=width / prs.slide_width,
                        h=height / prs.slide_height,
                    )
                    slots.append(
                        LayoutSlot(
                            placeholder_idx=placeholder.placeholder_format.idx,
                            placeholder_type=_map_placeholder_type(
                                placeholder.placeholder_format.type
                            ),
                            geometry=geometry,
                            normalized=normalized,
                            name=placeholder.name,
                        )
                    )
                existing_kinds = {slot.placeholder_type for slot in slots}
                inferred = _inferred_slots(layout, prs.slide_width, prs.slide_height)
                for slot in inferred:
                    if (
                        slot.placeholder_type not in existing_kinds
                        or slot.placeholder_type == PlaceholderType.BODY
                    ):
                        slots.append(slot)
                layout_type = classify_layout(layout.name, slots)
                layouts.append(
                    LayoutManifest(
                        layout_index=layout_index,
                        layout_name=layout.name,
                        layout_type=layout_type,
                        slots=slots,
                    )
                )
                layout_index_by_partname[str(layout.part.partname)] = layout_index
                layout_index += 1

        # Finished example slides are valuable design references even when their
        # layouts contain no formal placeholders.  Merge their text regions into
        # the owning layout as inferred slots; the builder will reuse the geometry
        # without copying example content into the generated deck.
        for slide_number, slide in enumerate(prs.slides, start=1):
            owner_index = layout_index_by_partname.get(
                str(slide.slide_layout.part.partname)
            )
            if owner_index is None:
                continue
            target = layouts[owner_index]
            existing = target.slots
            for slot in _inferred_slots(
                slide,
                prs.slide_width,
                prs.slide_height,
                slide_number * 1_000,
                include_text_content=True,
            ):
                overlaps = any(
                    abs(slot.normalized.x - current.normalized.x) < 0.025
                    and abs(slot.normalized.y - current.normalized.y) < 0.025
                    and abs(slot.normalized.w - current.normalized.w) < 0.04
                    and abs(slot.normalized.h - current.normalized.h) < 0.04
                    for current in existing
                )
                if not overlaps:
                    existing.append(slot)
            target.layout_type = classify_layout(target.layout_name, target.slots)

        manifest = TemplateManifest(
            source_hash=source_hash,
            slide_width_emu=prs.slide_width,
            slide_height_emu=prs.slide_height,
            colors=colors,
            fonts=fonts,
            brand_profile=brand_profile,
            layouts=layouts,
        )

        if cache_writable:
            try:
                cache_path.write_text(manifest.model_dump_json())
            except OSError:
                logger.warning(
                    "cannot write template cache %s; caching disabled", cache_path
                )
        return manifest
