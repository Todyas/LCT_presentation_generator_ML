from __future__ import annotations

from PIL import ImageFont
from pptx.text.text import TextFrame
from pptx.util import Pt

from app.models.template_manifest import Geometry

EMU_PER_INCH = 914_400


def _wrap_line_count(
    text: str, font: ImageFont.FreeTypeFont, max_width_px: float
) -> int:
    words = text.split() or [""]
    lines = 1
    line_width = 0.0
    space_width = font.getlength(" ")
    for word in words:
        word_width = font.getlength(word)
        if line_width == 0:
            line_width = word_width
        elif line_width + space_width + word_width <= max_width_px:
            line_width += space_width + word_width
        else:
            lines += 1
            line_width = word_width
    return lines


def autofit_font_size(
    paragraphs: list[str],
    box_width_emu: int,
    box_height_emu: int,
    font_path: str,
    min_size_pt: int = 10,
    max_size_pt: int = 44,
    line_spacing: float = 1.2,
    dpi: int = 96,
) -> int:
    box_width_px = box_width_emu / EMU_PER_INCH * dpi
    box_height_px = box_height_emu / EMU_PER_INCH * dpi

    def fits(size_pt: int) -> bool:
        font = ImageFont.truetype(font_path, int(size_pt * dpi / 72))
        # a single word wider than the box would be broken mid-word by the renderer
        if any(font.getlength(w) > box_width_px for p in paragraphs for w in p.split()):
            return False
        total_lines = sum(_wrap_line_count(p, font, box_width_px) for p in paragraphs)
        line_height_px = font.size * line_spacing
        return total_lines * line_height_px <= box_height_px

    if not fits(min_size_pt):
        return min_size_pt

    lo, hi, best = min_size_pt, max_size_pt, min_size_pt
    while lo <= hi:
        mid = (lo + hi) // 2
        if fits(mid):
            best = mid
            lo = mid + 1
        else:
            hi = mid - 1
    return best


def required_height_emu(
    paragraphs: list[str],
    box_width_emu: int,
    font_path: str,
    size_pt: int,
    line_spacing: float = 1.2,
    dpi: int = 96,
) -> int:
    """Height needed to render the paragraphs at exactly size_pt in a box of this width."""
    box_width_px = box_width_emu / EMU_PER_INCH * dpi
    font = ImageFont.truetype(font_path, int(size_pt * dpi / 72))
    total_lines = sum(_wrap_line_count(p, font, box_width_px) for p in paragraphs)
    return int(total_lines * font.size * line_spacing / dpi * EMU_PER_INCH)


def apply_autofit_to_text_frame(
    text_frame: TextFrame,
    geometry: Geometry,
    font_path: str,
    *,
    min_size_pt: int = 10,
    max_size_pt: int = 44,
    line_spacing: float = 1.2,
) -> None:
    paragraphs_text = [p.text for p in text_frame.paragraphs]
    size = autofit_font_size(
        paragraphs_text,
        box_width_emu=geometry.width_emu,
        box_height_emu=geometry.height_emu,
        font_path=font_path,
        min_size_pt=min_size_pt,
        max_size_pt=max_size_pt,
        line_spacing=line_spacing,
    )
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
