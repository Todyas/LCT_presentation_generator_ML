from __future__ import annotations

import colorsys
import io
import re
from dataclasses import dataclass
from functools import lru_cache

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_FILL_TYPE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.shapes.base import BaseShape
from pptx.shapes.placeholder import ChartPlaceholder, TablePlaceholder
from pptx.slide import Slide
from pptx.table import Table
from pptx.text.text import TextFrame
from pptx.util import Emu, Pt

from app.core.builder.autofit import (
    apply_autofit_to_text_frame,
    autofit_font_size,
    required_height_emu,
)
from app.models.presentation_ir import (
    BulletBlock,
    ChartData,
    ComparisonData,
    IconListData,
    MetricCard,
    ProcessData,
    TableData,
)
from app.models.template_manifest import Geometry, ThemeColors

_METRIC_VALUE_MAX_PT = 28
_METRIC_CARD_MIN_HEIGHT_EMU = 600_000
_BODY_MIN_PT = 14
_TITLE_MIN_PT = 24
_METRIC_LABEL_MAX_PT = 14
_METRIC_TEXT_MARGIN_PT = (
    4  # tight inner margin so long values like "2,4 млрд ₽" keep their width
)
_METRIC_CARD_GAP_EMU = (
    317_500  # 25pt gap between cards in a row, wide enough to stop edge clipping
)
_METRIC_CARD_MAX_HEIGHT_EMU = (
    1_600_200  # ~1.75in cap so a lone card "row" isn't absurdly tall
)

CHART_TYPE_MAP = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}


def _is_dark(hex_color: str) -> bool:
    red, green, blue = (int(hex_color[index : index + 2], 16) for index in (0, 2, 4))
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue < 118


def _surface_palette(
    slide: Slide, theme: ThemeColors, width: int, height: int
) -> tuple[str, str, str]:
    background = _background_hex(slide, theme, width, height)
    if _is_dark(background):
        surface = _mix_hex(background, "FFFFFF", 0.14)
        return surface, _contrasting_text_hex(surface, theme), theme.accent1
    surface = theme.lt1
    if _color_distance(surface, background) < 14:
        surface = (
            theme.lt2
            if _color_distance(theme.lt2, background) >= 14
            else _mix_hex(background, theme.dk1, 0.06)
        )
    return surface, _contrasting_text_hex(surface, theme), theme.accent1


def _style_text_frame(
    text_frame: TextFrame,
    text: str,
    color: str,
    font_name: str,
    size: int,
    *,
    bold: bool = False,
) -> None:
    text_frame.clear()
    text_frame.text = text
    text_frame.word_wrap = True
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = _rgb(color)


@dataclass(frozen=True)
class BBox:
    x: int
    y: int
    w: int
    h: int


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


_THEME_COLOR_MAP = {
    MSO_THEME_COLOR.DARK_1: "dk1",
    MSO_THEME_COLOR.DARK_2: "dk2",
    MSO_THEME_COLOR.LIGHT_1: "lt1",
    MSO_THEME_COLOR.LIGHT_2: "lt2",
    MSO_THEME_COLOR.TEXT_1: "dk1",
    MSO_THEME_COLOR.TEXT_2: "dk2",
    MSO_THEME_COLOR.BACKGROUND_1: "lt1",
    MSO_THEME_COLOR.BACKGROUND_2: "lt2",
    MSO_THEME_COLOR.ACCENT_1: "accent1",
    MSO_THEME_COLOR.ACCENT_2: "accent2",
    MSO_THEME_COLOR.ACCENT_3: "accent3",
    MSO_THEME_COLOR.ACCENT_4: "accent4",
    MSO_THEME_COLOR.ACCENT_5: "accent5",
    MSO_THEME_COLOR.ACCENT_6: "accent6",
}


def _color_hex(color, theme: ThemeColors) -> str | None:
    try:
        if color.rgb is not None:
            return str(color.rgb)
    except (AttributeError, ValueError):
        pass
    try:
        theme_key = _THEME_COLOR_MAP.get(color.theme_color)
    except (AttributeError, ValueError):
        theme_key = None
    return getattr(theme, theme_key) if theme_key is not None else None


_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_THEME_KEYS = frozenset(
    ("dk1", "dk2", "lt1", "lt2") + tuple(f"accent{index}" for index in range(1, 7))
)
_DEFAULT_CLR_MAP = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}


def _hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    return tuple(int(hex_color[index : index + 2], 16) for index in (0, 2, 4))  # type: ignore[return-value]


def _rgb_to_hex(red: float, green: float, blue: float) -> str:
    return "{:02X}{:02X}{:02X}".format(
        *(max(0, min(255, round(value))) for value in (red, green, blue))
    )


def _mix_hex(base: str, other: str, ratio: float) -> str:
    base_rgb, other_rgb = _hex_to_rgb(base), _hex_to_rgb(other)
    return _rgb_to_hex(
        *(b + (o - b) * ratio for b, o in zip(base_rgb, other_rgb, strict=True))
    )


def _color_distance(first: str, second: str) -> float:
    return max(
        abs(a - b) for a, b in zip(_hex_to_rgb(first), _hex_to_rgb(second), strict=True)
    )


def _relative_luminance(hex_color: str) -> float:
    channels = []
    for value in _hex_to_rgb(hex_color):
        srgb = value / 255
        channels.append(
            srgb / 12.92 if srgb <= 0.03928 else ((srgb + 0.055) / 1.055) ** 2.4
        )
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


def _contrast_ratio(first: str, second: str) -> float:
    lum_a, lum_b = _relative_luminance(first), _relative_luminance(second)
    lighter, darker = max(lum_a, lum_b), min(lum_a, lum_b)
    return (lighter + 0.05) / (darker + 0.05)


def _average_hex(colors: list[str]) -> str | None:
    if not colors:
        return None
    rgbs = [_hex_to_rgb(color) for color in colors]
    return _rgb_to_hex(
        *(sum(channel) / len(rgbs) for channel in zip(*rgbs, strict=True))
    )


@lru_cache(maxsize=256)
def _image_mean_hex(partname: str, blob: bytes) -> str | None:
    try:
        from PIL import Image, ImageStat

        with Image.open(io.BytesIO(blob)) as image:
            image = image.convert("RGB")
            image.thumbnail((64, 64))
            red, green, blue = ImageStat.Stat(image).mean[:3]
    except (OSError, ValueError, SyntaxError):
        return None
    return _rgb_to_hex(red, green, blue)


def _blip_mean_hexes(element, part) -> list[str]:
    colors = []
    for blip in element.iter(f"{{{_A_NS}}}blip"):
        rel_id = blip.get(f"{{{_R_NS}}}embed")
        if not rel_id:
            continue
        try:
            image_part = part.related_part(rel_id)
            mean = _image_mean_hex(str(image_part.partname), image_part.blob)
        except (KeyError, AttributeError, ValueError):
            continue
        if mean is not None:
            colors.append(mean)
    return colors


def _fill_hex(fill, theme: ThemeColors, part=None) -> str | None:
    fill_type = fill.type
    if fill_type == MSO_FILL_TYPE.SOLID:
        return _color_hex(fill.fore_color, theme)
    if fill_type == MSO_FILL_TYPE.GRADIENT:
        colors = [
            value
            for stop in fill.gradient_stops
            if (value := _color_hex(stop.color, theme)) is not None
        ]
        if colors:
            return colors[0]
    if fill_type == MSO_FILL_TYPE.PICTURE and part is not None:
        xpr = getattr(fill, "_xPr", None)
        if xpr is not None:
            return _average_hex(_blip_mean_hexes(xpr, part))
    return None


def _clr_map(master) -> dict[str, str]:
    mapping = dict(_DEFAULT_CLR_MAP)
    try:
        element = master._element.find(f"{{{_P_NS}}}clrMap")
    except AttributeError:
        return mapping
    if element is not None:
        for key in _DEFAULT_CLR_MAP:
            value = element.get(key)
            if value in _THEME_KEYS:
                mapping[key] = value
    return mapping


def _color_element_hex(
    element, theme: ThemeColors, clr_map: dict[str, str]
) -> str | None:
    tag = element.tag.rsplit("}", 1)[-1]
    if tag == "srgbClr":
        base = (element.get("val") or "").upper()
        if len(base) != 6:
            return None
    elif tag == "schemeClr":
        value = element.get("val") or ""
        key = clr_map.get(value, value)
        base = getattr(theme, key, None) if key in _THEME_KEYS else None
        if base is None:
            return None
    else:
        return None
    lum_mod = lum_off = None
    for child in element:
        name = child.tag.rsplit("}", 1)[-1]
        try:
            if name == "lumMod":
                lum_mod = int(child.get("val")) / 100_000
            elif name == "lumOff":
                lum_off = int(child.get("val")) / 100_000
        except (TypeError, ValueError):
            continue
    if lum_mod is None and lum_off is None:
        return base
    red, green, blue = (value / 255 for value in _hex_to_rgb(base))
    hue, lightness, saturation = colorsys.rgb_to_hls(red, green, blue)
    lightness = max(0.0, min(1.0, lightness * (lum_mod or 1.0) + (lum_off or 0.0)))
    return _rgb_to_hex(
        *(v * 255 for v in colorsys.hls_to_rgb(hue, lightness, saturation))
    )


def _bg_fill_element_hex(
    fill_element, part, theme: ThemeColors, clr_map: dict[str, str]
) -> str | None:
    tag = fill_element.tag.rsplit("}", 1)[-1]
    if tag == "solidFill":
        for child in fill_element:
            return _color_element_hex(child, theme, clr_map)
    elif tag == "gradFill":
        for stop in fill_element.iter(f"{{{_A_NS}}}gs"):
            for child in stop:
                color = _color_element_hex(child, theme, clr_map)
                if color is not None:
                    return color
    elif tag == "blipFill":
        return _average_hex(_blip_mean_hexes(fill_element, part))
    return None


def _owner_background_hex(
    owner, theme: ThemeColors, clr_map: dict[str, str]
) -> str | None:
    """Explicit background of a slide/layout/master, read without mutating it.

    python-pptx's ``.background.fill`` silently rewrites a missing or ``p:bgRef``
    background into ``noFill``, which would destroy the template's real look.
    """
    background = owner._element.find(f"{{{_P_NS}}}cSld/{{{_P_NS}}}bg")
    if background is None:
        return None
    properties = background.find(f"{{{_P_NS}}}bgPr")
    if properties is not None:
        for child in properties:
            color = _bg_fill_element_hex(child, owner.part, theme, clr_map)
            if color is not None:
                return color
        return None
    reference = background.find(f"{{{_P_NS}}}bgRef")
    if reference is not None:
        for child in reference:
            color = _color_element_hex(child, theme, clr_map)
            if color is not None:
                return color
    return None


def _covering_shape_hex(
    owner, theme: ThemeColors, slide_width_emu: int, slide_height_emu: int
) -> str | None:
    full_area = slide_width_emu * slide_height_emu
    candidates = []
    for shape in owner.shapes:
        try:
            width, height = shape.width, shape.height
        except (AttributeError, ValueError):
            continue
        if not width or not height or shape.is_placeholder:
            continue
        if (
            width * height >= full_area * 0.5
            and width >= slide_width_emu * 0.7
            and height >= slide_height_emu * 0.7
        ):
            candidates.append((width * height, shape))
    for _, shape in sorted(candidates, key=lambda item: item[0], reverse=True):
        color = _average_hex(_blip_mean_hexes(shape._element, owner.part))
        if color is None and hasattr(shape, "fill"):
            try:
                color = _fill_hex(shape.fill, theme, owner.part)
            except (AttributeError, TypeError, ValueError):
                color = None
        if color is not None:
            return color
    return None


def _background_hex(
    slide: Slide,
    theme: ThemeColors,
    slide_width_emu: int,
    slide_height_emu: int,
) -> str:
    layout = slide.slide_layout
    master = layout.slide_master
    clr_map = _clr_map(master)
    color = _owner_background_hex(slide, theme, clr_map)
    if color is not None:
        return color
    for owner in (layout, master):
        color = _covering_shape_hex(
            owner, theme, slide_width_emu, slide_height_emu
        ) or _owner_background_hex(owner, theme, clr_map)
        if color is not None:
            return color
    return theme.lt1


def _contrasting_text_hex(background_hex: str, theme: ThemeColors) -> str:
    best = max((theme.lt1, theme.dk1), key=lambda c: _contrast_ratio(c, background_hex))
    if _contrast_ratio(best, background_hex) >= 4.5:
        return best
    return "FFFFFF" if _relative_luminance(background_hex) < 0.18 else "000000"


def ensure_readable_text(
    slide: Slide, theme: ThemeColors, width: int, height: int
) -> None:
    """Give runs with no explicit colour one that contrasts with what is behind them."""
    slide_background = _background_hex(slide, theme, width, height)
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        background = None
        if hasattr(shape, "fill"):
            background = _fill_hex(shape.fill, theme)
        color = _contrasting_text_hex(background or slide_background, theme)
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.color.type is None:
                    run.font.color.rgb = _rgb(color)


_BOLD_MARKDOWN = re.compile(r"\*\*(.+?)\*\*")


def _render_bullet_text(paragraph, text: str) -> None:
    pos = 0
    for match in _BOLD_MARKDOWN.finditer(text):
        if match.start() > pos:
            paragraph.add_run().text = text[pos : match.start()]
        bold_run = paragraph.add_run()
        bold_run.text = match.group(1)
        bold_run.font.bold = True
        pos = match.end()
    if pos < len(text) or not paragraph.runs:
        paragraph.add_run().text = text[pos:]


def render_bullet_block(
    text_frame: TextFrame, block: BulletBlock, font_path: str, font_name: str
) -> None:
    text_frame.clear()
    for i, item in enumerate(block.items):
        paragraph = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        _render_bullet_text(paragraph, item.text)
        for run in paragraph.runs:
            run.font.name = font_name


def render_title_component(
    slide: Slide,
    geometry: Geometry,
    text: str,
    placeholder_shape: BaseShape | None,
    font_path: str,
    font_name: str,
    theme: ThemeColors,
    slide_width_emu: int,
    slide_height_emu: int,
    max_size_pt: float = 32,
) -> BBox:
    generated = placeholder_shape is None
    if placeholder_shape is None:
        shape = slide.shapes.add_textbox(
            Emu(geometry.left_emu),
            Emu(geometry.top_emu),
            Emu(geometry.width_emu),
            Emu(geometry.height_emu),
        )
    else:
        shape = placeholder_shape
    shape.text_frame.text = text
    shape.text_frame.word_wrap = True
    apply_autofit_to_text_frame(
        shape.text_frame,
        geometry,
        font_path,
        min_size_pt=_TITLE_MIN_PT,
        max_size_pt=max(int(max_size_pt), _TITLE_MIN_PT),
        line_spacing=1.15,
    )
    fallback_color = (
        _contrasting_text_hex(
            _background_hex(slide, theme, slide_width_emu, slide_height_emu), theme
        )
        if generated
        else None
    )
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.bold = True
            if fallback_color is not None:
                run.font.color.rgb = _rgb(fallback_color)
    return BBox(shape.left, shape.top, shape.width, shape.height)


def render_bullet_component(
    slide: Slide,
    geometry: Geometry,
    placeholder_shape: BaseShape | None,
    block: BulletBlock,
    font_path: str,
    font_name: str,
    theme: ThemeColors,
    slide_width_emu: int,
    slide_height_emu: int,
    max_size_pt: float = 24,
) -> BBox:
    generated_textbox = placeholder_shape is None
    target_shape = placeholder_shape
    if placeholder_shape is not None:
        text_frame = placeholder_shape.text_frame
        bbox = BBox(
            placeholder_shape.left,
            placeholder_shape.top,
            placeholder_shape.width,
            placeholder_shape.height,
        )
    else:
        textbox = slide.shapes.add_textbox(
            Emu(geometry.left_emu),
            Emu(geometry.top_emu),
            Emu(geometry.width_emu),
            Emu(geometry.height_emu),
        )
        target_shape = textbox
        text_frame = textbox.text_frame
        bbox = BBox(
            geometry.left_emu, geometry.top_emu, geometry.width_emu, geometry.height_emu
        )

    render_bullet_block(text_frame, block, font_path, font_name)
    if generated_textbox:
        text_frame.margin_left = Pt(6)
        text_frame.margin_right = Pt(6)
        text_frame.margin_top = Pt(6)
        text_frame.margin_bottom = Pt(6)
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    apply_autofit_to_text_frame(
        text_frame,
        geometry,
        font_path,
        min_size_pt=_BODY_MIN_PT,
        max_size_pt=max(int(max_size_pt), _BODY_MIN_PT),
        line_spacing=1.2,
    )
    fallback_color = (
        _contrasting_text_hex(
            _background_hex(slide, theme, slide_width_emu, slide_height_emu),
            theme,
        )
        if generated_textbox
        else None
    )
    for index, paragraph in enumerate(text_frame.paragraphs):
        paragraph.space_after = Pt(14)
        paragraph.line_spacing = 1.12
        if generated_textbox:
            paragraph.alignment = PP_ALIGN.LEFT
        paragraph_has_bold = any(run.font.bold for run in paragraph.runs)
        for run in paragraph.runs:
            if fallback_color is not None:
                run.font.color.rgb = _rgb(fallback_color)
            if index == 0 and not paragraph_has_bold:
                run.font.bold = True
    return _grow_to_fit_text(
        target_shape, text_frame, geometry, bbox, font_path, slide_height_emu
    )


def _grow_to_fit_text(
    shape: BaseShape,
    text_frame,
    geometry: Geometry,
    bbox: BBox,
    font_path: str,
    slide_height_emu: int,
) -> BBox:
    # at the font-size floor the text may still overflow: let the box grow downward
    # (never past the bottom safe margin) instead of shrinking the font further
    paragraphs = [p.text for p in text_frame.paragraphs]
    size = next(
        (
            int(run.font.size.pt)
            for p in text_frame.paragraphs
            for run in p.runs
            if run.font.size is not None
        ),
        _BODY_MIN_PT,
    )
    if size > _BODY_MIN_PT:
        return bbox
    inset = int(Pt(12))
    needed = (
        required_height_emu(
            paragraphs,
            max(bbox.w - 2 * inset, 1),
            font_path,
            size,
            line_spacing=1.2,
        )
        + len(paragraphs) * int(Pt(14))
        + 2 * inset
    )
    limit = int(slide_height_emu * 0.92) - bbox.y
    new_height = min(needed, limit)
    if new_height <= bbox.h:
        return bbox
    shape.left = Emu(bbox.x)
    shape.top = Emu(bbox.y)
    shape.width = Emu(bbox.w)
    shape.height = Emu(new_height)
    return BBox(bbox.x, bbox.y, bbox.w, new_height)


def render_metric_card(
    slide: Slide,
    geometry: Geometry,
    card: MetricCard,
    theme: ThemeColors,
    font_path: str,
    font_name: str | None = None,
) -> BBox:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Emu(geometry.left_emu),
        Emu(geometry.top_emu),
        Emu(geometry.width_emu),
        Emu(geometry.height_emu),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(theme.lt2)
    shape.line.color.rgb = _rgb(theme.accent1)

    margin_emu = int(Pt(_METRIC_TEXT_MARGIN_PT))
    usable_width = geometry.width_emu - 2 * margin_emu

    value_height = int(geometry.height_emu * 0.6)
    value_box = slide.shapes.add_textbox(
        Emu(geometry.left_emu),
        Emu(geometry.top_emu),
        Emu(geometry.width_emu),
        Emu(value_height),
    )
    value_box.text_frame.margin_left = Pt(_METRIC_TEXT_MARGIN_PT)
    value_box.text_frame.margin_right = Pt(_METRIC_TEXT_MARGIN_PT)
    value_box.text_frame.text = card.value
    value_size = autofit_font_size(
        [card.value],
        usable_width,
        value_height,
        font_path,
        max_size_pt=_METRIC_VALUE_MAX_PT,
    )
    value_box.text_frame.paragraphs[0].runs[0].font.size = Pt(value_size)
    value_box.text_frame.paragraphs[0].runs[0].font.bold = True
    value_box.text_frame.paragraphs[0].runs[0].font.color.rgb = _rgb(theme.accent1)
    if font_name:
        value_box.text_frame.paragraphs[0].runs[0].font.name = font_name

    label_top = geometry.top_emu + value_height
    label_height = geometry.height_emu - value_height
    label_box = slide.shapes.add_textbox(
        Emu(geometry.left_emu),
        Emu(label_top),
        Emu(geometry.width_emu),
        Emu(label_height),
    )
    label_box.text_frame.margin_left = Pt(_METRIC_TEXT_MARGIN_PT)
    label_box.text_frame.margin_right = Pt(_METRIC_TEXT_MARGIN_PT)
    label = card.label if not card.delta else f"{card.label} · {card.delta}"
    label_box.text_frame.text = label
    label_size = autofit_font_size(
        [label],
        usable_width,
        label_height,
        font_path,
        min_size_pt=12,
        max_size_pt=_METRIC_LABEL_MAX_PT,
    )
    label_box.text_frame.paragraphs[0].runs[0].font.size = Pt(label_size)
    label_box.text_frame.paragraphs[0].runs[0].font.color.rgb = _rgb(theme.dk1)
    if font_name:
        label_box.text_frame.paragraphs[0].runs[0].font.name = font_name

    return BBox(
        geometry.left_emu, geometry.top_emu, geometry.width_emu, geometry.height_emu
    )


def render_metric_card_group(
    slide: Slide,
    geometry: Geometry,
    cards: list[MetricCard],
    theme: ThemeColors,
    font_path: str,
    font_name: str | None = None,
    slide_width_emu: int | None = None,
) -> list[BBox]:
    n = len(cards)
    card_width = (geometry.width_emu - _METRIC_CARD_GAP_EMU * (n - 1)) // n
    if slide_width_emu is not None and card_width < slide_width_emu * 0.18 and n > 1:
        # too narrow side by side: stack the cards in a column instead
        gap = _METRIC_CARD_GAP_EMU // 2
        card_height = max(
            (geometry.height_emu - gap * (n - 1)) // n, _METRIC_CARD_MIN_HEIGHT_EMU
        )
        card_height = min(card_height, _METRIC_CARD_MAX_HEIGHT_EMU)
        return [
            render_metric_card(
                slide,
                Geometry(
                    left_emu=geometry.left_emu,
                    top_emu=geometry.top_emu + i * (card_height + gap),
                    width_emu=geometry.width_emu,
                    height_emu=card_height,
                ),
                card,
                theme,
                font_path,
                font_name,
            )
            for i, card in enumerate(cards)
        ]
    card_height = min(geometry.height_emu, _METRIC_CARD_MAX_HEIGHT_EMU)

    boxes: list[BBox] = []
    for i, card in enumerate(cards):
        card_geometry = Geometry(
            left_emu=geometry.left_emu + i * (card_width + _METRIC_CARD_GAP_EMU),
            top_emu=geometry.top_emu,
            width_emu=card_width,
            height_emu=card_height,
        )
        boxes.append(
            render_metric_card(slide, card_geometry, card, theme, font_path, font_name)
        )
    return boxes


def fill_table_cells(
    pptx_table: Table,
    table: TableData,
    theme: ThemeColors | None = None,
    font_name: str | None = None,
) -> None:
    for c, header in enumerate(table.headers):
        pptx_table.cell(0, c).text = header
    for r, row in enumerate(table.rows, start=1):
        for c, cell_value in enumerate(row):
            pptx_table.cell(r, c).text = cell_value
    if theme is None:
        return
    for c in range(len(table.headers)):
        cell = pptx_table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(theme.accent1)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = _rgb(theme.lt1)
                if font_name:
                    run.font.name = font_name
    for r in range(1, len(table.rows) + 1):
        for c in range(len(table.headers)):
            cell = pptx_table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(theme.lt1 if r % 2 else theme.lt2)
            cell.margin_left = Pt(8)
            cell.margin_right = Pt(8)
            cell.margin_top = Pt(5)
            cell.margin_bottom = Pt(5)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.color.rgb = _rgb(theme.dk1)
                    run.font.size = Pt(14 if len(table.rows) < 6 else 12)
                    run.font.bold = c == 0
                    if font_name:
                        run.font.name = font_name


def render_table(
    slide: Slide,
    geometry: Geometry,
    table: TableData,
    theme: ThemeColors | None = None,
    font_name: str | None = None,
) -> BBox:
    rows = len(table.rows) + 1
    cols = len(table.headers)
    graphic_frame = slide.shapes.add_table(
        rows,
        cols,
        Emu(geometry.left_emu),
        Emu(geometry.top_emu),
        Emu(geometry.width_emu),
        Emu(geometry.height_emu),
    )
    fill_table_cells(graphic_frame.table, table, theme, font_name)
    return BBox(
        geometry.left_emu, geometry.top_emu, geometry.width_emu, geometry.height_emu
    )


def render_table_into_placeholder(
    placeholder: TablePlaceholder,
    table: TableData,
    theme: ThemeColors | None = None,
    font_name: str | None = None,
) -> BBox:
    graphic_frame = placeholder.insert_table(len(table.rows) + 1, len(table.headers))
    fill_table_cells(graphic_frame.table, table, theme, font_name)
    return BBox(
        graphic_frame.left, graphic_frame.top, graphic_frame.width, graphic_frame.height
    )


def render_table_component(
    slide: Slide,
    geometry: Geometry,
    placeholder_shape: BaseShape | None,
    table: TableData,
    theme: ThemeColors,
    font_name: str | None = None,
    font_path: str = "",
    slide_width_emu: int = 12_192_000,
    slide_height_emu: int = 6_858_000,
) -> BBox:
    if isinstance(placeholder_shape, TablePlaceholder):
        return render_table_into_placeholder(placeholder_shape, table, theme, font_name)
    return render_table(slide, geometry, table, theme, font_name)


def build_chart_data(chart: ChartData) -> CategoryChartData:
    chart_data = CategoryChartData()
    chart_data.categories = chart.categories
    for series in chart.series:
        chart_data.add_series(series.name, series.values)
    return chart_data


def _style_chart(chart_shape, theme: ThemeColors | None) -> None:
    chart_shape.chart_style = 2
    chart_shape.has_legend = len(chart_shape.series) > 1
    if theme is None:
        return
    try:
        chart_shape.chart_area.format.fill.background()
        chart_shape.plot_area.format.fill.background()
    except (AttributeError, ValueError):
        pass
    palette = [theme.accent1, theme.accent2, theme.accent5, theme.accent6]
    for index, series in enumerate(chart_shape.series):
        color = _rgb(palette[index % len(palette)])
        try:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
        except (AttributeError, ValueError):
            pass
        try:
            series.format.line.color.rgb = color
        except (AttributeError, ValueError):
            pass
    for axis_name in ("category_axis", "value_axis"):
        try:
            axis = getattr(chart_shape, axis_name)
            axis.tick_labels.font.name = "Arial"
            axis.tick_labels.font.size = Pt(10)
            axis.tick_labels.font.color.rgb = _rgb(theme.dk1)
            axis.format.line.color.rgb = _rgb(theme.lt2)
            if axis_name == "value_axis":
                axis.major_gridlines.format.line.color.rgb = _rgb(theme.lt2)
        except (AttributeError, ValueError):
            pass
    try:
        chart_shape.legend.font.size = Pt(10)
        chart_shape.legend.font.color.rgb = _rgb(theme.dk1)
    except (AttributeError, ValueError):
        pass


def render_chart(
    slide: Slide, geometry: Geometry, chart: ChartData, theme: ThemeColors | None = None
) -> BBox:
    if theme is not None:
        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(geometry.left_emu),
            Emu(geometry.top_emu),
            Emu(geometry.width_emu),
            Emu(geometry.height_emu),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = _rgb(theme.lt1)
        panel.line.color.rgb = _rgb(theme.lt2)
    chart_data = build_chart_data(chart)
    inset = 100_000 if theme is not None else 0
    graphic_frame = slide.shapes.add_chart(
        CHART_TYPE_MAP[chart.chart_type],
        Emu(geometry.left_emu + inset),
        Emu(geometry.top_emu + inset),
        Emu(geometry.width_emu - 2 * inset),
        Emu(geometry.height_emu - 2 * inset),
        chart_data,
    )
    _style_chart(graphic_frame.chart, theme)
    return BBox(
        geometry.left_emu, geometry.top_emu, geometry.width_emu, geometry.height_emu
    )


def render_chart_into_placeholder(
    placeholder: ChartPlaceholder, chart: ChartData, theme: ThemeColors | None = None
) -> BBox:
    graphic_frame = placeholder.insert_chart(
        CHART_TYPE_MAP[chart.chart_type], build_chart_data(chart)
    )
    _style_chart(graphic_frame.chart, theme)
    return BBox(
        graphic_frame.left, graphic_frame.top, graphic_frame.width, graphic_frame.height
    )


def render_chart_component(
    slide: Slide,
    geometry: Geometry,
    placeholder_shape: BaseShape | None,
    chart: ChartData,
    theme: ThemeColors | None = None,
) -> BBox:
    return render_chart(slide, geometry, chart, theme)


def render_comparison(
    slide: Slide,
    geometry: Geometry,
    comparison: ComparisonData,
    theme: ThemeColors,
    font_name: str,
    slide_width_emu: int,
    slide_height_emu: int,
    font_path: str = "",
) -> BBox:
    surface, text_color, accent = _surface_palette(
        slide, theme, slide_width_emu, slide_height_emu
    )
    gap = 180_000
    width = (geometry.width_emu - gap) // 2
    sides = (
        (comparison.left_title, comparison.left_items, theme.accent2),
        (comparison.right_title, comparison.right_items, accent),
    )
    body_size = _BODY_MIN_PT
    for index, (heading, items, side_color) in enumerate(sides):
        left = geometry.left_emu + index * (width + gap)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(left),
            Emu(geometry.top_emu),
            Emu(width),
            Emu(geometry.height_emu),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(surface)
        card.line.color.rgb = _rgb(side_color)
        band_height = min(650_000, geometry.height_emu // 4)
        band = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(left),
            Emu(geometry.top_emu),
            Emu(width),
            Emu(band_height),
        )
        band.fill.solid()
        band.fill.fore_color.rgb = _rgb(side_color)
        band.line.fill.background()
        _style_text_frame(
            band.text_frame,
            heading,
            _contrasting_text_hex(side_color, theme),
            font_name,
            18,
            bold=True,
        )
        band.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        band.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        body = slide.shapes.add_textbox(
            Emu(left + 90_000),
            Emu(geometry.top_emu + band_height + 70_000),
            Emu(width - 180_000),
            Emu(geometry.height_emu - band_height - 140_000),
        )
        block = BulletBlock(items=[{"text": value} for value in items])
        render_bullet_block(body.text_frame, block, "", font_name)
        for paragraph in body.text_frame.paragraphs:
            paragraph.space_after = Pt(10)
            for run in paragraph.runs:
                run.font.color.rgb = _rgb(text_color)
                run.font.size = Pt(body_size)
        if font_path:
            apply_autofit_to_text_frame(
                body.text_frame,
                Geometry(
                    left_emu=left + 90_000,
                    top_emu=geometry.top_emu + band_height + 70_000,
                    width_emu=width - 180_000,
                    height_emu=geometry.height_emu - band_height - 140_000,
                ),
                font_path,
                min_size_pt=_BODY_MIN_PT,
                max_size_pt=body_size,
                line_spacing=1.08,
            )
    return BBox(
        geometry.left_emu, geometry.top_emu, geometry.width_emu, geometry.height_emu
    )


def render_process(
    slide: Slide,
    geometry: Geometry,
    process: ProcessData,
    theme: ThemeColors,
    font_name: str,
    slide_width_emu: int,
    slide_height_emu: int,
    font_path: str = "",
) -> BBox:
    surface, text_color, accent = _surface_palette(
        slide, theme, slide_width_emu, slide_height_emu
    )
    count = len(process.steps)
    gap_x, gap_y = 150_000, 130_000
    per_row_cap = 4 if geometry.width_emu < slide_width_emu * 0.6 else 6
    fit_cap = max(1, (geometry.width_emu + gap_x) // int(slide_width_emu * 0.18 + gap_x))
    per_row_cap = max(1, min(per_row_cap, fit_cap))
    rows = max(1, -(-count // per_row_cap))
    columns = max(1, -(-count // rows))
    step_width = (geometry.width_emu - gap_x * (columns - 1)) // columns
    step_height = (geometry.height_emu - gap_y * (rows - 1)) // rows
    for index, step in enumerate(process.steps):
        row, column = divmod(index, columns)
        left = geometry.left_emu + column * (step_width + gap_x)
        top = geometry.top_emu + row * (step_height + gap_y)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(left),
            Emu(top),
            Emu(step_width),
            Emu(step_height),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(surface)
        card.line.color.rgb = _rgb(
            theme.dk2 if not _is_dark(surface) else theme.accent1
        )
        band = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(left),
            Emu(top),
            Emu(58_000),
            Emu(step_height),
        )
        band.fill.solid()
        band.fill.fore_color.rgb = _rgb(accent)
        band.line.fill.background()
        badge_size = min(390_000, step_height // 3)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(left + 150_000),
            Emu(top + 120_000),
            Emu(badge_size),
            Emu(badge_size),
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = _rgb(accent)
        badge.line.fill.background()
        _style_text_frame(
            badge.text_frame, str(index + 1), theme.lt1, font_name, 16, bold=True
        )
        badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        title_box = slide.shapes.add_textbox(
            Emu(left + badge_size + 230_000),
            Emu(top + 105_000),
            Emu(step_width - badge_size - 320_000),
            Emu(badge_size + 50_000),
        )
        _style_text_frame(
            title_box.text_frame, step.title, text_color, font_name, 14, bold=True
        )
        description = slide.shapes.add_textbox(
            Emu(left + 150_000),
            Emu(top + badge_size + 190_000),
            Emu(step_width - 270_000),
            Emu(max(100_000, step_height - badge_size - 280_000)),
        )
        _style_text_frame(
            description.text_frame,
            step.description or step.title,
            text_color,
            font_name,
            _BODY_MIN_PT,
        )
        if font_path:
            apply_autofit_to_text_frame(
                title_box.text_frame,
                Geometry(
                    left_emu=left,
                    top_emu=top,
                    width_emu=step_width - badge_size - 250_000,
                    height_emu=badge_size,
                ),
                font_path,
                min_size_pt=_BODY_MIN_PT,
                max_size_pt=16,
                line_spacing=1.0,
            )
            apply_autofit_to_text_frame(
                description.text_frame,
                Geometry(
                    left_emu=left,
                    top_emu=top,
                    width_emu=step_width - 270_000,
                    height_emu=max(100_000, step_height - badge_size - 280_000),
                ),
                font_path,
                min_size_pt=12,
                max_size_pt=_BODY_MIN_PT,
                line_spacing=1.05,
            )
    return BBox(
        geometry.left_emu, geometry.top_emu, geometry.width_emu, geometry.height_emu
    )


def render_icon_list(
    slide: Slide,
    geometry: Geometry,
    icon_list: IconListData,
    theme: ThemeColors,
    font_name: str,
    slide_width_emu: int,
    slide_height_emu: int,
    font_path: str = "",
) -> BBox:
    surface, text_color, accent = _surface_palette(
        slide, theme, slide_width_emu, slide_height_emu
    )
    gap = 120_000
    columns = min(len(icon_list.items), 2 if len(icon_list.items) <= 4 else 3)
    while columns > 1 and (
        (geometry.width_emu - gap * (columns - 1)) // columns < slide_width_emu * 0.2
    ):
        columns -= 1
    rows = (len(icon_list.items) + columns - 1) // columns
    cell_width = (geometry.width_emu - gap * (columns - 1)) // columns
    cell_height = (geometry.height_emu - gap * (rows - 1)) // rows
    icon_size = min(390_000, cell_height - 150_000)
    for index, item in enumerate(icon_list.items):
        row, column = divmod(index, columns)
        left = geometry.left_emu + column * (cell_width + gap)
        top = geometry.top_emu + row * (cell_height + gap)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(left),
            Emu(top),
            Emu(cell_width),
            Emu(cell_height),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(surface)
        card.line.color.rgb = _rgb(
            theme.dk2 if not _is_dark(surface) else theme.accent1
        )
        icon = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(left + 80_000),
            Emu(top + (cell_height - icon_size) // 2),
            Emu(icon_size),
            Emu(icon_size),
        )
        icon.fill.solid()
        icon.fill.fore_color.rgb = _rgb(accent)
        icon.line.fill.background()
        # Small brand bars are deliberately geometric and consistent; using
        # Unicode symbols made the deck depend on fallback fonts and look like
        # stock SmartArt.
        for bar_index, ratio in enumerate((0.42, 0.68, 0.88)):
            bar = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Emu(left + 145_000),
                Emu(top + (cell_height - icon_size) // 2 + 90_000 + bar_index * 85_000),
                Emu(int((icon_size - 130_000) * ratio)),
                Emu(36_000),
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = _rgb(theme.lt1)
            bar.line.fill.background()
        text_box = slide.shapes.add_textbox(
            Emu(left + icon_size + 150_000),
            Emu(top + 60_000),
            Emu(cell_width - icon_size - 220_000),
            Emu(cell_height - 120_000),
        )
        _style_text_frame(
            text_box.text_frame,
            item.title if not item.description else f"{item.title}\n{item.description}",
            text_color,
            font_name,
            _BODY_MIN_PT,
        )
        text_box.text_frame.paragraphs[0].runs[0].font.bold = True
        text_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        if font_path:
            apply_autofit_to_text_frame(
                text_box.text_frame,
                Geometry(
                    left_emu=left,
                    top_emu=top,
                    width_emu=cell_width - icon_size - 220_000,
                    height_emu=cell_height - 120_000,
                ),
                font_path,
                min_size_pt=12,
                max_size_pt=_BODY_MIN_PT,
                line_spacing=1.05,
            )
    return BBox(
        geometry.left_emu, geometry.top_emu, geometry.width_emu, geometry.height_emu
    )


def render_image_placeholder(
    slide: Slide, geometry: Geometry, alt_text: str, theme: ThemeColors
) -> BBox:
    """Render a brand-colored vector illustration when no external asset exists."""
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Emu(geometry.left_emu),
        Emu(geometry.top_emu),
        Emu(geometry.width_emu),
        Emu(geometry.height_emu),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = _rgb(theme.dk2)
    panel.line.color.rgb = _rgb(theme.accent5)
    size = min(geometry.width_emu, geometry.height_emu) // 3
    centers = (
        (0.22, 0.3, theme.accent1),
        (0.58, 0.2, theme.accent5),
        (0.5, 0.58, theme.accent2),
    )
    for x_ratio, y_ratio, color in centers:
        node = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Emu(geometry.left_emu + int(geometry.width_emu * x_ratio)),
            Emu(geometry.top_emu + int(geometry.height_emu * y_ratio)),
            Emu(size),
            Emu(size),
        )
        node.fill.solid()
        node.fill.fore_color.rgb = _rgb(color)
        node.line.fill.background()
    # Keep the semantic description in the file without showing placeholder copy.
    try:
        panel.element.nvSpPr.cNvPr.set("descr", alt_text)
    except AttributeError:
        pass
    return BBox(
        geometry.left_emu, geometry.top_emu, geometry.width_emu, geometry.height_emu
    )
