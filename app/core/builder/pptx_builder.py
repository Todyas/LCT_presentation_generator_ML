from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.presentation import Presentation as PptxPresentation
from pptx.shapes.base import BaseShape
from pptx.slide import Slide, SlideLayout

from app.core.builder.shape_factory import (
    BBox,
    ensure_readable_text,
    render_bullet_component,
    render_chart_component,
    render_comparison,
    render_icon_list,
    render_image_placeholder,
    render_metric_card_group,
    render_process,
    render_table_component,
    render_title_component,
)
from app.core.parser.font_resolver import resolve_font_path
from app.core.parser.template_parser import _PLACEHOLDER_TYPE_MAP
from app.models.presentation_ir import (
    BulletBlock,
    ChartData,
    ComparisonData,
    IconListData,
    ImagePlaceholder,
    MetricCard,
    PresentationIR,
    ProcessData,
    SlideComponent,
    SlideIR,
    TableData,
)
from app.models.template_manifest import (
    Geometry,
    LayoutManifest,
    PlaceholderType,
    TemplateManifest,
)


@dataclass
class BuildResult:
    pptx_path: str
    bbox_map: dict[int, dict[str, BBox]] = field(default_factory=dict)


SLIDE_MARGIN_EMU = (
    91_440  # 0.1 inch — spacing between an auto-placed component and its predecessor
)
_FALLBACK_WIDTH_RATIO = (
    0.8  # fraction of slide width used only when a layout has no content slots at all
)
_COLUMN_GAP_EMU = (
    137_160  # 0.15in gap between two bullet-block columns sharing one slot
)

# invert Sprint 1's PP_PLACEHOLDER -> PlaceholderType map so lookups stay consistent with the parser
_OUR_TO_PPTX_PLACEHOLDER_TYPES: dict[PlaceholderType, set[int]] = {}
for _pptx_type, _our_type in _PLACEHOLDER_TYPE_MAP.items():
    _OUR_TO_PPTX_PLACEHOLDER_TYPES.setdefault(_our_type, set()).add(int(_pptx_type))

_COMPONENT_TO_PLACEHOLDER_TYPE: dict[type, PlaceholderType] = {
    BulletBlock: PlaceholderType.BODY,
    TableData: PlaceholderType.TABLE,
    ChartData: PlaceholderType.CHART,
    ImagePlaceholder: PlaceholderType.PICTURE,
    ComparisonData: PlaceholderType.BODY,
    ProcessData: PlaceholderType.BODY,
    IconListData: PlaceholderType.BODY,
}

# footer/date/slide-number chrome usually spans the full canvas width regardless of the
# layout's real content-safe region, so it must not seed the fallback bounding box (Bug 3)
_CHROME_PLACEHOLDER_TYPES = frozenset(
    {
        PlaceholderType.TITLE,
        PlaceholderType.FOOTER,
        PlaceholderType.DATE,
        PlaceholderType.SLIDE_NUMBER,
    }
)


_MIN_INFERRED_W = 0.30
_MIN_INFERRED_H = 0.15
_MIN_EXPLICIT_W = 0.18
_MIN_EXPLICIT_H = 0.10
_MIN_COLUMN_RATIO = 0.18
_MIN_GEOMETRY_H_EMU = 400_000


def _slot_usable(slot, layout: LayoutManifest) -> bool:
    n = slot.normalized
    if not slot.inferred:
        return n.w >= _MIN_EXPLICIT_W and n.h >= _MIN_EXPLICIT_H
    if n.w < _MIN_INFERRED_W or n.h < _MIN_INFERRED_H:
        return False
    region = layout.content_region
    if region is None:
        return True
    g = slot.geometry
    iw = min(g.right_emu, region.right_emu) - max(g.left_emu, region.left_emu)
    ih = min(g.bottom_emu, region.bottom_emu) - max(g.top_emu, region.top_emu)
    if iw <= 0 or ih <= 0:
        return False
    return iw * ih >= 0.6 * g.width_emu * g.height_emu


def _group_weight(group: list[SlideComponent]) -> float:
    if isinstance(group[0], MetricCard):
        return 0.45
    if isinstance(group[0], ImagePlaceholder):
        return 0.8
    return 1.0


def _region_geometry(
    layout: LayoutManifest, cursor_bottom_emu: int, share: float, remaining: int
) -> Geometry | None:
    region = layout.content_region
    if region is None:
        return None
    top = region.top_emu
    if (
        cursor_bottom_emu > top
        and region.bottom_emu - cursor_bottom_emu >= 0.2 * region.height_emu
    ):
        top = cursor_bottom_emu
    available = region.bottom_emu - top - SLIDE_MARGIN_EMU * (remaining - 1)
    height = min(int(available * share), region.bottom_emu - top)
    if height <= 0:
        return None
    return Geometry(
        left_emu=region.left_emu,
        top_emu=top,
        width_emu=region.width_emu,
        height_emu=max(height, min(_MIN_GEOMETRY_H_EMU, region.bottom_emu - top)),
    )


def _clear_title(geometry: Geometry, title_box: BBox | None) -> Geometry:
    if title_box is None:
        return geometry
    overlaps_x = (
        geometry.left_emu < title_box.x + title_box.w
        and geometry.right_emu > title_box.x
    )
    overlaps_y = (
        geometry.top_emu < title_box.y + title_box.h
        and geometry.bottom_emu > title_box.y
    )
    if not (overlaps_x and overlaps_y):
        return geometry
    top = title_box.y + title_box.h + SLIDE_MARGIN_EMU
    height = max(geometry.bottom_emu - top, _MIN_GEOMETRY_H_EMU)
    return Geometry(
        left_emu=geometry.left_emu,
        top_emu=top,
        width_emu=geometry.width_emu,
        height_emu=height,
    )


def _group_components(
    components: list[SlideComponent], pair_bullet_blocks: bool
) -> list[list[SlideComponent]]:
    groups: list[list[SlideComponent]] = []
    i = 0
    while i < len(components):
        component = components[i]
        if isinstance(component, MetricCard):
            j = i
            while j < len(components) and isinstance(components[j], MetricCard):
                j += 1
            groups.append(components[i:j])
            i = j
            continue
        if (
            pair_bullet_blocks
            and isinstance(component, BulletBlock)
            and i + 1 < len(components)
            and isinstance(components[i + 1], BulletBlock)
        ):
            groups.append([component, components[i + 1]])
            i += 2
            continue
        groups.append([component])
        i += 1
    return groups


def _layout_for_slide(manifest: TemplateManifest, slide_ir: SlideIR) -> LayoutManifest:
    if slide_ir.layout_index is not None:
        for candidate in manifest.layouts:
            if candidate.layout_index == slide_ir.layout_index:
                return candidate
    return manifest.find_layout_or_fallback(slide_ir.layout_type, slide_ir.slide_index)


class PptxBuilder:
    def build(
        self, template_path: str, manifest: TemplateManifest, ir: PresentationIR
    ) -> BuildResult:
        prs = Presentation(template_path)
        self._strip_existing_slides(prs)
        all_layouts = self._all_slide_layouts(prs)

        bbox_map: dict[int, dict[str, BBox]] = {}
        for slide_ir in ir.slides:
            layout = _layout_for_slide(manifest, slide_ir)
            pptx_layout = all_layouts[layout.layout_index]
            slide = prs.slides.add_slide(pptx_layout)

            shape_boxes: dict[str, BBox] = {}
            font_path = resolve_font_path(manifest.fonts.minor_latin)

            title_placeholder = self._find_placeholder(slide, PlaceholderType.TITLE)
            title_slot = layout.slot_by_type(PlaceholderType.TITLE)
            used_placeholder_idxs: set[int] = set()
            if title_placeholder is not None or title_slot is not None:
                title_geometry = (
                    title_slot.geometry
                    if title_slot is not None
                    else Geometry(
                        left_emu=title_placeholder.left,
                        top_emu=title_placeholder.top,
                        width_emu=title_placeholder.width,
                        height_emu=title_placeholder.height,
                    )
                )
                shape_boxes["title"] = render_title_component(
                    slide,
                    title_geometry,
                    slide_ir.title.text,
                    title_placeholder,
                    font_path,
                    manifest.fonts.major_latin,
                    manifest.colors,
                    manifest.slide_width_emu,
                    manifest.slide_height_emu,
                    manifest.brand_profile.title_size_pt,
                )
                if title_placeholder is not None:
                    used_placeholder_idxs.add(title_placeholder.placeholder_format.idx)
                elif title_slot is not None:
                    used_placeholder_idxs.add(title_slot.placeholder_idx)
            else:
                title_geometry = Geometry(
                    left_emu=int(prs.slide_width * 0.07),
                    top_emu=int(prs.slide_height * 0.055),
                    width_emu=int(prs.slide_width * 0.86),
                    height_emu=int(prs.slide_height * 0.17),
                )
                shape_boxes["title"] = render_title_component(
                    slide,
                    title_geometry,
                    slide_ir.title.text,
                    None,
                    font_path,
                    manifest.fonts.major_latin,
                    manifest.colors,
                    manifest.slide_width_emu,
                    manifest.slide_height_emu,
                    manifest.brand_profile.title_size_pt,
                )

            cursor_bottom_emu = (
                shape_boxes["title"].y + shape_boxes["title"].h
                if "title" in shape_boxes
                else 0
            )

            bullet_block_count = sum(
                1 for c in slide_ir.components if isinstance(c, BulletBlock)
            )
            available_body_slots = sum(
                1
                for s in layout.slots
                if s.placeholder_type == PlaceholderType.BODY
                and s.placeholder_idx not in used_placeholder_idxs
                and _slot_usable(s, layout)
            )
            pair_bullet_blocks = bullet_block_count == 2 and available_body_slots < 2

            component_index = 0
            title_box = shape_boxes.get("title")
            groups = _group_components(slide_ir.components, pair_bullet_blocks)
            weights = [_group_weight(g) for g in groups]
            for group_index, group in enumerate(groups):
                region_kwargs = {
                    "title_box": title_box,
                    "remaining": len(groups) - group_index,
                    "share": weights[group_index] / sum(weights[group_index:]),
                }
                if isinstance(group[0], MetricCard):
                    geometry, slot_used, placeholder_shape = self._resolve_geometry(
                        slide,
                        group[0],
                        layout,
                        used_placeholder_idxs,
                        prs,
                        cursor_bottom_emu,
                        **region_kwargs,
                    )
                    if slot_used is not None:
                        used_placeholder_idxs.add(slot_used)
                    if placeholder_shape is not None:
                        self._remove_shape(placeholder_shape)
                    boxes = render_metric_card_group(
                        slide,
                        geometry,
                        group,
                        manifest.colors,
                        font_path,
                        manifest.fonts.minor_latin,
                        slide_width_emu=manifest.slide_width_emu,
                    )
                    for bbox in boxes:
                        shape_boxes[f"component_{component_index}"] = bbox
                        component_index += 1
                    cursor_bottom_emu = (
                        geometry.top_emu + geometry.height_emu + SLIDE_MARGIN_EMU
                    )
                    continue

                if len(group) == 2 and all(isinstance(c, BulletBlock) for c in group):
                    # the matched slot (or fallback box) becomes a shared container split into two
                    # columns, so the original single-region placeholder is left unclaimed here and
                    # gets purged below rather than showing its own ghost text under the columns
                    container_geometry, _, _ = self._resolve_geometry(
                        slide,
                        group[0],
                        layout,
                        used_placeholder_idxs,
                        prs,
                        cursor_bottom_emu,
                        **region_kwargs,
                    )
                    for block, col_geometry in zip(
                        group,
                        self._split_into_columns(
                            container_geometry, manifest.slide_width_emu
                        ),
                    ):
                        bbox = render_bullet_component(
                            slide,
                            col_geometry,
                            None,
                            block,
                            font_path,
                            manifest.fonts.minor_latin,
                            manifest.colors,
                            manifest.slide_width_emu,
                            manifest.slide_height_emu,
                            manifest.brand_profile.body_size_pt,
                        )
                        shape_boxes[f"component_{component_index}"] = bbox
                        component_index += 1
                    cursor_bottom_emu = (
                        container_geometry.top_emu
                        + container_geometry.height_emu
                        + SLIDE_MARGIN_EMU
                    )
                    continue

                component = group[0]
                geometry, slot_used, placeholder_shape = self._resolve_geometry(
                    slide,
                    component,
                    layout,
                    used_placeholder_idxs,
                    prs,
                    cursor_bottom_emu,
                    **region_kwargs,
                )
                if slot_used is not None:
                    used_placeholder_idxs.add(slot_used)

                bbox = self._render_component(
                    slide, component, geometry, placeholder_shape, manifest, font_path
                )
                shape_boxes[f"component_{component_index}"] = bbox
                component_index += 1
                cursor_bottom_emu = bbox.y + bbox.h + SLIDE_MARGIN_EMU

            self._purge_unused_placeholders(slide, used_placeholder_idxs)
            ensure_readable_text(
                slide,
                manifest.colors,
                manifest.slide_width_emu,
                manifest.slide_height_emu,
            )
            bbox_map[slide_ir.slide_index] = shape_boxes

        output_path = str(Path(template_path).with_name(f"built_{ir.variant}.pptx"))
        prs.save(output_path)
        return BuildResult(pptx_path=output_path, bbox_map=bbox_map)

    def _all_slide_layouts(self, prs: PptxPresentation) -> list[SlideLayout]:
        return [
            layout for master in prs.slide_masters for layout in master.slide_layouts
        ]

    def _split_into_columns(
        self, geometry: Geometry, slide_width_emu: int | None = None
    ) -> tuple[Geometry, Geometry]:
        col_width = (geometry.width_emu - _COLUMN_GAP_EMU) // 2
        if (
            slide_width_emu is not None
            and col_width < slide_width_emu * _MIN_COLUMN_RATIO
        ):
            half = (geometry.height_emu - _COLUMN_GAP_EMU) // 2
            top_part = Geometry(
                left_emu=geometry.left_emu,
                top_emu=geometry.top_emu,
                width_emu=geometry.width_emu,
                height_emu=half,
            )
            bottom_part = Geometry(
                left_emu=geometry.left_emu,
                top_emu=geometry.top_emu + half + _COLUMN_GAP_EMU,
                width_emu=geometry.width_emu,
                height_emu=half,
            )
            return top_part, bottom_part
        left_col = Geometry(
            left_emu=geometry.left_emu,
            top_emu=geometry.top_emu,
            width_emu=col_width,
            height_emu=geometry.height_emu,
        )
        right_col = Geometry(
            left_emu=geometry.left_emu + col_width + _COLUMN_GAP_EMU,
            top_emu=geometry.top_emu,
            width_emu=col_width,
            height_emu=geometry.height_emu,
        )
        return left_col, right_col

    def _strip_existing_slides(self, prs: PptxPresentation) -> None:
        xml_slides = prs.slides._sldIdLst
        slide_ids = list(xml_slides)
        for slide_id in slide_ids:
            rId = slide_id.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            prs.part.drop_rel(rId)
            xml_slides.remove(slide_id)

    def _find_placeholder(
        self, slide: Slide, ptype: PlaceholderType
    ) -> BaseShape | None:
        target = _OUR_TO_PPTX_PLACEHOLDER_TYPES.get(ptype, set())
        for shape in slide.placeholders:
            if shape.placeholder_format.type in target:
                return shape
        return None

    def _find_placeholder_by_idx(self, slide: Slide, idx: int) -> BaseShape | None:
        for shape in slide.placeholders:
            if shape.placeholder_format.idx == idx:
                return shape
        return None

    def _remove_shape(self, shape: BaseShape) -> None:
        shape.element.getparent().remove(shape.element)

    def _purge_unused_placeholders(
        self, slide: Slide, used_placeholder_idxs: set[int]
    ) -> None:
        for shape in list(slide.placeholders):
            if shape.placeholder_format.idx not in used_placeholder_idxs:
                self._remove_shape(shape)

    def _resolve_geometry(
        self,
        slide: Slide,
        component: SlideComponent,
        layout: LayoutManifest,
        used_placeholder_idxs: set[int],
        prs: PptxPresentation,
        cursor_bottom_emu: int,
        title_box: BBox | None = None,
        share: float = 1.0,
        remaining: int = 1,
    ) -> tuple[Geometry, int | None, BaseShape | None]:
        from_slots = self._resolve_from_slots(
            slide, component, layout, used_placeholder_idxs, prs, cursor_bottom_emu
        )
        if from_slots is not None:
            geometry, slot_idx, shape = from_slots
            slot = next(
                (s for s in layout.slots if s.placeholder_idx == slot_idx), None
            )
            if slot is None or slot.inferred:
                geometry = _clear_title(geometry, title_box)
            return geometry, slot_idx, shape

        claimed = self._claim_unusable_body_slot(
            slide, layout, used_placeholder_idxs
        )
        region_geometry = _region_geometry(layout, cursor_bottom_emu, share, remaining)
        if region_geometry is not None:
            return _clear_title(region_geometry, title_box), claimed[0], claimed[1]
        if isinstance(
            component,
            (ComparisonData, ProcessData, IconListData, ChartData, TableData),
        ):
            geometry = self._compute_visual_geometry(layout, prs, cursor_bottom_emu)
        else:
            geometry = self._compute_fallback_geometry(layout, prs, cursor_bottom_emu)
        return _clear_title(geometry, title_box), claimed[0], claimed[1]

    def _claim_unusable_body_slot(
        self,
        slide: Slide,
        layout: LayoutManifest,
        used_placeholder_idxs: set[int],
    ) -> tuple[int | None, BaseShape | None]:
        # claim an explicit body placeholder we chose not to use so the caller can
        # drop it and its prompt text does not show through
        for slot in layout.slots:
            if (
                slot.placeholder_type == PlaceholderType.BODY
                and not slot.inferred
                and slot.placeholder_idx not in used_placeholder_idxs
            ):
                shape = self._find_placeholder_by_idx(slide, slot.placeholder_idx)
                if shape is not None:
                    return slot.placeholder_idx, shape
        return None, None

    def _resolve_from_slots(
        self,
        slide: Slide,
        component: SlideComponent,
        layout: LayoutManifest,
        used_placeholder_idxs: set[int],
        prs: PptxPresentation,
        cursor_bottom_emu: int,
    ) -> tuple[Geometry, int | None, BaseShape | None] | None:
        wanted_type = _COMPONENT_TO_PLACEHOLDER_TYPE.get(type(component))
        visual_component = isinstance(
            component,
            (ComparisonData, ProcessData, IconListData, ChartData, TableData),
        )
        free_slots = [
            slot
            for slot in layout.slots
            if slot.placeholder_idx not in used_placeholder_idxs
            and _slot_usable(slot, layout)
        ]
        if visual_component:
            matching_slots = [
                s for s in free_slots if s.placeholder_type == wanted_type
            ]
            # Real wide chart/table placeholders are useful. Narrow inferred
            # sample boxes are not: they caused five-step diagrams to be
            # squeezed into the left third of a slide.
            if matching_slots:
                slot = matching_slots[0]
                if slot.geometry.width_emu >= int(prs.slide_width * 0.58):
                    return (
                        slot.geometry,
                        slot.placeholder_idx,
                        self._find_placeholder_by_idx(slide, slot.placeholder_idx),
                    )
            body_slots = [
                s for s in free_slots if s.placeholder_type == PlaceholderType.BODY
            ]
            explicit_body_slots = [slot for slot in body_slots if not slot.inferred]
            if explicit_body_slots:
                first = explicit_body_slots[0]
                left = min(slot.geometry.left_emu for slot in explicit_body_slots)
                top = min(slot.geometry.top_emu for slot in explicit_body_slots)
                right = max(slot.geometry.right_emu for slot in explicit_body_slots)
                bottom = max(slot.geometry.bottom_emu for slot in explicit_body_slots)
                return (
                    Geometry(
                        left_emu=left,
                        top_emu=top,
                        width_emu=right - left,
                        height_emu=bottom - top,
                    ),
                    first.placeholder_idx,
                    self._find_placeholder_by_idx(slide, first.placeholder_idx),
                )
            return None
        if wanted_type is not None:
            for slot in free_slots:
                if slot.placeholder_type == wanted_type:
                    placeholder_shape = self._find_placeholder_by_idx(
                        slide, slot.placeholder_idx
                    )
                    return slot.geometry, slot.placeholder_idx, placeholder_shape

        # A native chart/table/card can be drawn inside a generic BODY placeholder.
        # Reusing that full content region preserves the template's safe margins and
        # avoids stacking a shallow fallback box immediately below the title.
        body_slots = [
            s for s in free_slots if s.placeholder_type == PlaceholderType.BODY
        ]
        if body_slots:
            first = body_slots[0]
            placeholder_shape = self._find_placeholder_by_idx(
                slide, first.placeholder_idx
            )
            if wanted_type != PlaceholderType.BODY and len(body_slots) > 1:
                left = min(slot.geometry.left_emu for slot in body_slots)
                top = min(slot.geometry.top_emu for slot in body_slots)
                right = max(slot.geometry.right_emu for slot in body_slots)
                bottom = max(slot.geometry.bottom_emu for slot in body_slots)
                return (
                    Geometry(
                        left_emu=left,
                        top_emu=top,
                        width_emu=right - left,
                        height_emu=bottom - top,
                    ),
                    first.placeholder_idx,
                    placeholder_shape,
                )
            return first.geometry, first.placeholder_idx, placeholder_shape
        return None

    def _compute_visual_geometry(
        self, layout: LayoutManifest, prs: PptxPresentation, cursor_bottom_emu: int
    ) -> Geometry:
        """Large predictable canvas for charts and custom visual components."""
        margin_x = int(prs.slide_width * 0.07)
        bottom = int(prs.slide_height * 0.90)
        top = max(cursor_bottom_emu + SLIDE_MARGIN_EMU, int(prs.slide_height * 0.24))
        if top >= bottom - int(prs.slide_height * 0.30):
            top = int(prs.slide_height * 0.25)
        return Geometry(
            left_emu=margin_x,
            top_emu=top,
            width_emu=int(prs.slide_width * 0.86),
            height_emu=max(400_000, bottom - top),
        )

    def _compute_fallback_geometry(
        self, layout: LayoutManifest, prs: PptxPresentation, cursor_bottom_emu: int
    ) -> Geometry:
        content_slots = [
            s
            for s in layout.slots
            if s.placeholder_type not in _CHROME_PLACEHOLDER_TYPES
            and _slot_usable(s, layout)
        ]
        if content_slots:
            left = min(s.geometry.left_emu for s in content_slots)
            right = max(s.geometry.right_emu for s in content_slots)
            width = right - left
            region_top = min(s.geometry.top_emu for s in content_slots)
            region_bottom = max(s.geometry.bottom_emu for s in content_slots)
            top = max(cursor_bottom_emu, region_top)
            available_height = region_bottom - top
        else:
            width = int(prs.slide_width * _FALLBACK_WIDTH_RATIO)
            left = int((prs.slide_width - width) / 2)
            top = cursor_bottom_emu
            available_height = prs.slide_height - top - int(prs.slide_height * 0.07)
        fallback_height = max(400_000, available_height)
        return Geometry(
            left_emu=left,
            top_emu=top,
            width_emu=width,
            height_emu=fallback_height,
        )

    def _render_component(
        self,
        slide: Slide,
        component: SlideComponent,
        geometry: Geometry,
        placeholder_shape: BaseShape | None,
        manifest: TemplateManifest,
        font_path: str,
    ) -> BBox:
        if isinstance(component, BulletBlock):
            return render_bullet_component(
                slide,
                geometry,
                placeholder_shape,
                component,
                font_path,
                manifest.fonts.minor_latin,
                manifest.colors,
                manifest.slide_width_emu,
                manifest.slide_height_emu,
                manifest.brand_profile.body_size_pt,
            )
        if isinstance(component, TableData):
            bbox = render_table_component(
                slide,
                geometry,
                placeholder_shape,
                component,
                manifest.colors,
                manifest.fonts.minor_latin,
                font_path,
                manifest.slide_width_emu,
                manifest.slide_height_emu,
            )
            if placeholder_shape is not None and not hasattr(
                placeholder_shape, "insert_table"
            ):
                self._remove_shape(placeholder_shape)
            return bbox
        if isinstance(component, ChartData):
            bbox = render_chart_component(
                slide, geometry, placeholder_shape, component, manifest.colors
            )
            if placeholder_shape is not None:
                self._remove_shape(placeholder_shape)
            return bbox
        if isinstance(component, ComparisonData):
            if placeholder_shape is not None:
                self._remove_shape(placeholder_shape)
            return render_comparison(
                slide,
                geometry,
                component,
                manifest.colors,
                manifest.fonts.minor_latin,
                manifest.slide_width_emu,
                manifest.slide_height_emu,
                font_path,
            )
        if isinstance(component, ProcessData):
            if placeholder_shape is not None:
                self._remove_shape(placeholder_shape)
            return render_process(
                slide,
                geometry,
                component,
                manifest.colors,
                manifest.fonts.minor_latin,
                manifest.slide_width_emu,
                manifest.slide_height_emu,
                font_path,
            )
        if isinstance(component, IconListData):
            if placeholder_shape is not None:
                self._remove_shape(placeholder_shape)
            return render_icon_list(
                slide,
                geometry,
                component,
                manifest.colors,
                manifest.fonts.minor_latin,
                manifest.slide_width_emu,
                manifest.slide_height_emu,
                font_path,
            )
        if isinstance(component, ImagePlaceholder):
            bbox = render_image_placeholder(
                slide, geometry, component.alt_text, manifest.colors
            )
            if placeholder_shape is not None:
                self._remove_shape(placeholder_shape)
            return bbox
        raise TypeError(f"unhandled component type: {type(component)!r}")
