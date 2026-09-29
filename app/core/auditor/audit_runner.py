from __future__ import annotations

from collections.abc import Iterator

from pptx import Presentation

from app.core.agents.llm_client import LLMClient
from app.core.agents.prompt_registry import PromptRegistry
from app.core.auditor.contrast_audit import run_contrast_audit
from app.core.auditor.density_audit import (
    run_density_audit_on_ir,
    run_placeholder_text_audit,
)
from app.core.auditor.geometry_audit import BBox, run_geometry_audit
from app.core.auditor.semantic_audit import run_semantic_audit
from app.core.auditor.visual_audit import run_visual_variety_audit
from app.core.builder.pptx_builder import BuildResult
from app.core.builder.shape_factory import _background_hex, _clr_map, _fill_hex
from app.models.audit_report import AuditReport
from app.models.presentation_ir import PresentationIR
from app.models.template_manifest import TemplateManifest


async def run_full_audit(
    build_result: BuildResult,
    ir: PresentationIR,
    manifest: TemplateManifest,
    brief: str,
    llm: LLMClient,
    registry: PromptRegistry,
    model: str,
) -> AuditReport:
    issues = []

    for slide_index, shape_boxes in build_result.bbox_map.items():
        geometry_bboxes = {
            shape_id: BBox(x=box.x, y=box.y, w=box.w, h=box.h)
            for shape_id, box in shape_boxes.items()
        }
        issues.extend(
            run_geometry_audit(
                geometry_bboxes,
                manifest.slide_width_emu,
                manifest.slide_height_emu,
                slide_index,
            )
        )

    for (
        slide_index,
        shape_id,
        hex_color,
        background,
        is_large,
    ) in _iter_text_shape_colors(build_result, ir, manifest):
        issue = run_contrast_audit(
            text_color_hex=hex_color,
            bg_color_hex=background,
            is_large_text=is_large,
            slide_index=slide_index,
            shape_id=shape_id,
        )
        if issue is not None:
            issues.append(issue)

    issues.extend(run_density_audit_on_ir(ir))
    issues.extend(run_visual_variety_audit(ir))
    issues.extend(run_placeholder_text_audit(build_result.pptx_path))
    issues.extend(await run_semantic_audit(ir, brief, llm, registry, model))

    return AuditReport(variant=ir.variant, issues=issues)


def _iter_text_shape_colors(
    build_result: BuildResult, ir: PresentationIR, manifest: TemplateManifest
) -> Iterator[tuple[int, str, str, str, bool]]:
    prs = Presentation(build_result.pptx_path)
    for slide_index, slide in enumerate(prs.slides):
        default_color = getattr(
            manifest.colors, _clr_map(slide.slide_layout.slide_master)["tx1"]
        )
        slide_background = _background_hex(
            slide,
            manifest.colors,
            manifest.slide_width_emu,
            manifest.slide_height_emu,
        )
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            shape_background = (
                _fill_hex(shape.fill, manifest.colors)
                if hasattr(shape, "fill")
                else None
            ) or slide_background
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if not run.text:
                        continue
                    color = run.font.color
                    if color.type is None:
                        hex_color = default_color
                    else:
                        try:
                            hex_color = str(color.rgb)
                        except AttributeError:
                            hex_color = default_color
                    size_pt = run.font.size.pt if run.font.size is not None else 12
                    is_large = size_pt >= (14 if run.font.bold else 18)
                    yield (
                        slide_index,
                        f"slide{slide_index}_shape{shape.shape_id}",
                        hex_color,
                        shape_background,
                        is_large,
                    )
