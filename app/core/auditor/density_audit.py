from __future__ import annotations

import re

from pptx import Presentation

from app.models.audit_report import AuditIssue, IssueType, Severity
from app.models.presentation_ir import (
    MAX_BULLETS,
    MAX_TABLE_COLS,
    MAX_TABLE_ROWS,
    MAX_WORDS_PER_BULLET,
    PresentationIR,
)

_PLACEHOLDER_PATTERN = re.compile(
    r"\b(TODO|Lorem ipsum|XXX|заглушка|ТБД)\b", re.IGNORECASE
)


def run_density_audit_on_ir(ir: PresentationIR) -> list[AuditIssue]:
    issues: list[AuditIssue] = []
    for slide in ir.slides:
        for component in slide.components:
            if component.type == "bullet_block":
                if len(component.items) > MAX_BULLETS:
                    issues.append(
                        AuditIssue(
                            issue_type=IssueType.DENSITY_BULLETS,
                            severity=Severity.CRITICAL,
                            slide_index=slide.slide_index,
                            message=f"{len(component.items)} bullets exceeds max {MAX_BULLETS}",
                            auto_fixable=False,
                        )
                    )
                for item in component.items:
                    word_count = len(item.text.split())
                    if word_count > MAX_WORDS_PER_BULLET:
                        issues.append(
                            AuditIssue(
                                issue_type=IssueType.DENSITY_WORDS,
                                severity=Severity.CRITICAL,
                                slide_index=slide.slide_index,
                                message=f"bullet with {word_count} words exceeds max {MAX_WORDS_PER_BULLET}: {item.text!r}",
                                auto_fixable=False,
                            )
                        )
            if component.type == "table" and (
                len(component.headers) > MAX_TABLE_COLS or len(component.rows) > MAX_TABLE_ROWS
            ):
                issues.append(
                    AuditIssue(
                        issue_type=IssueType.TABLE_SIZE,
                        severity=Severity.CRITICAL,
                        slide_index=slide.slide_index,
                        message=f"table {len(component.rows)}x{len(component.headers)} exceeds max {MAX_TABLE_ROWS}x{MAX_TABLE_COLS}",
                        auto_fixable=False,
                    )
                )
    return issues


def run_placeholder_text_audit(pptx_path: str) -> list[AuditIssue]:
    issues: list[AuditIssue] = []
    prs = Presentation(pptx_path)
    for slide_index, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text
            match = _PLACEHOLDER_PATTERN.search(text)
            if match:
                issues.append(
                    AuditIssue(
                        issue_type=IssueType.PLACEHOLDER_TEXT,
                        severity=Severity.CRITICAL,
                        slide_index=slide_index,
                        message=f"placeholder text {match.group(0)!r} found in shape text: {text!r}",
                        shape_ids=[str(shape.shape_id)],
                        auto_fixable=False,
                    )
                )
    return issues
