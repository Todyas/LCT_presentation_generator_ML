from __future__ import annotations

import asyncio
import io
import json
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Annotated

from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    Form,
    HTTPException,
    UploadFile,
)
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from pptx import Presentation
from redis.asyncio import Redis
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.schemas import (
    ActivateRevisionRequest,
    HealthResponse,
    JobCreatedResponse,
    JobStatusResponse,
    JobSummaryDTO,
    SlideRevisionRequest,
    VariantResultDTO,
)
from app.config import Settings, get_settings
from app.core.parser.template_parser import TemplateParser
from app.pipeline.jobs import JobStatus, JobStore, JobType, SlideRevisionStore
from app.pipeline.service import run_generation_job, run_revision_job

router = APIRouter()
_settings = get_settings()
job_store = JobStore(
    _settings.database_url, initialize_schema=not _settings.task_queue_enabled
)
revision_store = SlideRevisionStore(
    _settings.database_url, initialize_schema=not _settings.task_queue_enabled
)

_VARIANT_META = {
    "A": {
        "name": "Executive",
        "audience": "Для руководства",
        "description": "Крупные выводы и KPI, минимум текста.",
    },
    "B": {
        "name": "Analytical",
        "audience": "Для проектной защиты",
        "description": "Больше данных, сравнений, таблиц и графиков.",
    },
    "C": {
        "name": "Pitch",
        "audience": "Для выступления",
        "description": "Крупные визуальные блоки и короткие сообщения для рассказа со сцены.",
    },
}


def _validate_pptx_upload(raw_bytes: bytes, max_bytes: int) -> None:
    if len(raw_bytes) > max_bytes:
        raise HTTPException(
            413, f"template exceeds max upload size of {max_bytes} bytes"
        )
    try:
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
            if "ppt/presentation.xml" not in zf.namelist():
                raise HTTPException(
                    422, "not a valid .pptx file (missing ppt/presentation.xml)"
                )
    except zipfile.BadZipFile:
        raise HTTPException(422, "not a valid .pptx file (not a zip archive)")


@router.get("/health", response_model=HealthResponse)
async def health() -> HealthResponse:
    return HealthResponse()


@router.get("/health/ready", response_model=HealthResponse)
async def readiness(
    settings: Annotated[Settings, Depends(get_settings)],
) -> HealthResponse:
    try:
        with Session(job_store.engine) as session:
            session.execute(text("SELECT 1"))
        if settings.task_queue_enabled:
            redis = Redis.from_url(settings.redis_url)
            try:
                await redis.ping()
            finally:
                await redis.aclose()
    except Exception as exc:
        raise HTTPException(503, f"service dependencies are not ready: {exc}") from exc
    return HealthResponse()


def _enqueue_job(
    *,
    job_id: str,
    revision: bool,
    background_tasks: BackgroundTasks,
    settings: Settings,
) -> None:
    if settings.task_queue_enabled:
        try:
            if revision:
                from app.worker.tasks import revise_slide_task

                queued = revise_slide_task.delay(job_id)
            else:
                from app.worker.tasks import generate_deck_task

                queued = generate_deck_task.delay(job_id)
        except Exception as exc:
            job_store.update(
                job_id,
                status=JobStatus.FAILED,
                stage="queue_failed",
                progress=100,
                error=str(exc),
            )
            if revision:
                child = job_store.get(job_id)
                parent = job_store.get(child.parent_job_id) if child else None
                if (
                    child is not None
                    and parent is not None
                    and parent.result is not None
                    and child.variant in parent.result.variants
                ):
                    parent.result.variants[child.variant].export_state = "READY"
                    job_store.update(parent.job_id, result=parent.result)
            raise HTTPException(503, "background queue is unavailable") from exc
        job_store.update(job_id, task_id=queued.id)
        return
    runner = run_revision_job if revision else run_generation_job
    background_tasks.add_task(runner, job_id, settings)


@router.post("/templates/analyze")
async def analyze_template(
    settings: Annotated[Settings, Depends(get_settings)],
    template: Annotated[UploadFile, File()],
) -> JSONResponse:
    """Return the Template DNA needed by the upload screen before generation."""

    raw_bytes = await template.read()
    _validate_pptx_upload(raw_bytes, settings.max_upload_bytes)
    presentation = Presentation(io.BytesIO(raw_bytes))

    temp_path: str | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as temp_file:
            temp_file.write(raw_bytes)
            temp_path = temp_file.name
        manifest = TemplateParser().parse(temp_path)
    finally:
        if temp_path is not None:
            Path(temp_path).unlink(missing_ok=True)

    known_layouts = [
        layout for layout in manifest.layouts if layout.layout_type.value != "UNKNOWN"
    ]
    match_score = round(100 * len(known_layouts) / max(1, len(manifest.layouts)))
    colors = list(dict.fromkeys(manifest.colors.model_dump().values()))
    fonts = list(dict.fromkeys(manifest.fonts.model_dump().values()))
    layout_types = sorted({layout.layout_type.value for layout in known_layouts})
    return JSONResponse(
        content={
            "filename": template.filename or "template.pptx",
            "source_hash": manifest.source_hash,
            "status": "parsed",
            "match_score": match_score,
            "slide_count": len(presentation.slides),
            "layout_count": len(manifest.layouts),
            "master_count": len(presentation.slide_masters),
            "colors": colors,
            "fonts": fonts,
            "layout_types": layout_types,
        }
    )


@router.post("/generate", response_model=JobCreatedResponse, status_code=202)
async def generate(
    background_tasks: BackgroundTasks,
    settings: Annotated[Settings, Depends(get_settings)],
    template: Annotated[UploadFile, File()],
    brief: Annotated[str, Form(min_length=10, max_length=5000)],
    slide_count: Annotated[int, Form(ge=10, le=15)] = 12,
    purpose: Annotated[str, Form()] = "project",
    language: Annotated[str, Form()] = "ru",
    style: Annotated[str, Form()] = "balanced",
) -> JobCreatedResponse:
    raw_bytes = await template.read()
    _validate_pptx_upload(raw_bytes, settings.max_upload_bytes)

    job = job_store.create(
        template_filename=template.filename or "template.pptx",
        brief=brief,
        purpose=purpose,
        slide_count=slide_count,
        language=language,
        style=style,
    )
    storage_dir = Path(settings.storage_dir) / job.job_id
    storage_dir.mkdir(parents=True, exist_ok=True)
    template_path = storage_dir / "template.pptx"
    template_path.write_bytes(raw_bytes)
    job_store.update(job.job_id, template_path=str(template_path))

    _enqueue_job(
        job_id=job.job_id,
        revision=False,
        background_tasks=background_tasks,
        settings=settings,
    )
    return JobCreatedResponse(job_id=job.job_id)


@router.get("/jobs", response_model=list[JobSummaryDTO])
async def list_jobs() -> list[JobSummaryDTO]:
    return [
        JobSummaryDTO(
            job_id=j.job_id,
            status=j.status.value,
            created_at=j.created_at,
            template_filename=j.template_filename,
            job_type=j.job_type.value,
            parent_job_id=j.parent_job_id,
        )
        for j in job_store.list_all()
    ]


@router.get("/jobs/{job_id}/events")
async def stream_job_events(job_id: str) -> StreamingResponse:
    if job_store.get(job_id) is None:
        raise HTTPException(404, "job not found")

    async def event_stream():
        last_payload = None
        while True:
            job = job_store.get(job_id)
            if job is None:
                yield 'event: deleted\ndata: {"status":"DELETED"}\n\n'
                return
            payload = json.dumps(
                {
                    "job_id": job.job_id,
                    "status": job.status.value,
                    "stage": job.stage,
                    "progress": job.progress,
                    "error": job.error,
                },
                ensure_ascii=False,
            )
            if payload != last_payload:
                yield f"data: {payload}\n\n"
                last_payload = payload
            if job.status in {JobStatus.DONE, JobStatus.PARTIAL, JobStatus.FAILED}:
                return
            await asyncio.sleep(0.5)

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@router.delete("/jobs/{job_id}", status_code=204)
async def delete_job(
    job_id: str, settings: Annotated[Settings, Depends(get_settings)]
) -> Response:
    job = job_store.get(job_id)
    if job is None:
        raise HTTPException(404, "job not found")
    if job.status == JobStatus.RUNNING:
        raise HTTPException(409, "a running job cannot be deleted")
    job_store.delete(job_id)
    storage_dir = Path(settings.storage_dir) / job_id
    shutil.rmtree(storage_dir, ignore_errors=True)
    return Response(status_code=204)


@router.get("/jobs/{job_id}", response_model=JobStatusResponse)
async def get_job(job_id: str) -> JobStatusResponse:
    job = job_store.get(job_id)
    if job is None:
        raise HTTPException(404, "job not found")

    variants: list[VariantResultDTO] = []
    if job.result is not None:
        for v in job.result.variants.values():
            variants.append(
                VariantResultDTO(
                    variant=v.variant,
                    pptx_available=v.pptx_path is not None,
                    pdf_available=v.pdf_path is not None,
                    html_available=v.html_path is not None,
                    preview_count=len(v.preview_paths),
                    audit_passed=v.audit_report.passed if v.audit_report else None,
                    error=v.error,
                    revision=v.revision,
                    export_state=v.export_state,
                )
            )

    return JobStatusResponse(
        job_id=job.job_id,
        status=job.status.value,
        variants=variants,
        error=job.error,
        stage=job.stage,
        progress=job.progress,
        job_type=job.job_type.value,
        parent_job_id=job.parent_job_id,
        output=job.output,
    )


@router.get("/jobs/{job_id}/files/{variant}/{kind}")
async def get_job_file(job_id: str, variant: str, kind: str) -> FileResponse:
    job = job_store.get(job_id)
    if job is None or job.result is None:
        raise HTTPException(404, "job not found or not finished")
    variant_result = job.result.variants.get(variant)
    if variant_result is None:
        raise HTTPException(404, "variant not found")
    if variant_result.export_state == "STALE":
        raise HTTPException(409, "exports are being rebuilt for this variant")

    path_by_kind = {
        "pptx": variant_result.pptx_path,
        "pdf": variant_result.pdf_path,
        "html": variant_result.html_path,
    }
    if kind not in path_by_kind or path_by_kind[kind] is None:
        raise HTTPException(
            404, f"file kind {kind!r} not available for variant {variant!r}"
        )
    return FileResponse(path_by_kind[kind])


@router.get("/jobs/{job_id}/previews/{variant}/{slide_position}")
async def get_slide_preview(
    job_id: str, variant: str, slide_position: int
) -> FileResponse:
    job = job_store.get(job_id)
    if job is None or job.result is None:
        raise HTTPException(404, "job not found or not finished")
    variant_result = job.result.variants.get(variant.upper())
    if variant_result is None:
        raise HTTPException(404, "variant not found")
    if slide_position < 1 or slide_position > len(variant_result.preview_paths):
        raise HTTPException(404, "slide preview not found")
    return FileResponse(
        variant_result.preview_paths[slide_position - 1], media_type="image/png"
    )


def _variant_payload(job_id: str, variant_result, manifest=None) -> dict:
    code = variant_result.variant
    meta = _VARIANT_META.get(code, {"name": code, "audience": "", "description": ""})
    ir = variant_result.presentation_ir
    issues = variant_result.audit_report.issues if variant_result.audit_report else []
    slides = []
    if ir is not None:
        for position, slide in enumerate(ir.slides, start=1):
            slide_issues = []
            for issue in issues:
                if issue.slide_index != slide.slide_index:
                    continue
                payload = issue.model_dump(mode="json")
                if issue.bbox is not None and manifest is not None:
                    payload["bbox_normalized"] = {
                        "x": issue.bbox.x / manifest.slide_width_emu,
                        "y": issue.bbox.y / manifest.slide_height_emu,
                        "w": issue.bbox.w / manifest.slide_width_emu,
                        "h": issue.bbox.h / manifest.slide_height_emu,
                    }
                slide_issues.append(payload)
            severities = {item["severity"] for item in slide_issues}
            audit_status = (
                "critical"
                if "CRITICAL" in severities
                else "warning"
                if "WARNING" in severities
                else "ok"
            )
            slides.append(
                {
                    "position": position,
                    "slide_index": slide.slide_index,
                    "title": slide.title.text,
                    "layout_type": slide.layout_type.value,
                    "preview_url": (
                        f"/jobs/{job_id}/previews/{code}/{position}"
                        if position <= len(variant_result.preview_paths)
                        else None
                    ),
                    "audit_status": audit_status,
                    "issues": slide_issues,
                }
            )
    critical_count = sum(issue.severity.value == "CRITICAL" for issue in issues)
    warning_count = sum(issue.severity.value == "WARNING" for issue in issues)
    audit_score = max(0, 100 - critical_count * 15 - warning_count * 5)
    text_words = 0
    action_titles = 0
    data_components = 0
    component_count = 0
    if ir is not None:
        for slide in ir.slides:
            text_words += len(slide.title.text.split())
            action_titles += int(slide.title.is_action_title)
            for component in slide.components:
                component_count += 1
                if component.type == "bullet_block":
                    text_words += sum(
                        len(item.text.split()) for item in component.items
                    )
                if component.type in {"metric_card", "chart", "table"}:
                    data_components += 1
    slide_total = len(ir.slides) if ir is not None else 0
    metrics = {
        "text_density": min(100, round(100 * text_words / max(1, slide_total * 60))),
        "conclusions": round(100 * action_titles / max(1, slide_total)),
        "data": round(100 * data_components / max(1, component_count)),
    }
    return {
        "code": code,
        **meta,
        "revision": variant_result.revision,
        "export_state": variant_result.export_state,
        "status": "failed" if variant_result.error else "ready",
        "error": variant_result.error,
        "audit_score": audit_score,
        "metrics": metrics,
        "slides": slides,
        "exports": {
            "pptx": f"/jobs/{job_id}/files/{code}/pptx"
            if variant_result.pptx_path and variant_result.export_state == "READY"
            else None,
            "pdf": f"/jobs/{job_id}/files/{code}/pdf"
            if variant_result.pdf_path and variant_result.export_state == "READY"
            else None,
            "html": f"/jobs/{job_id}/files/{code}/html"
            if variant_result.html_path and variant_result.export_state == "READY"
            else None,
        },
    }


@router.get("/jobs/{job_id}/result")
async def get_job_result(job_id: str) -> JSONResponse:
    job = job_store.get(job_id)
    if job is None or job.result is None:
        raise HTTPException(404, "job not found or not finished")
    manifest = job.result.template_manifest
    template_dna = None
    if manifest is not None:
        known_layouts = [
            layout
            for layout in manifest.layouts
            if layout.layout_type.value != "UNKNOWN"
        ]
        template_dna = {
            "source_hash": manifest.source_hash,
            "colors": list(dict.fromkeys(manifest.colors.model_dump().values())),
            "fonts": list(dict.fromkeys(manifest.fonts.model_dump().values())),
            "layout_count": len(manifest.layouts),
            "match_score": round(
                100 * len(known_layouts) / max(1, len(manifest.layouts))
            ),
            "slide_width_emu": manifest.slide_width_emu,
            "slide_height_emu": manifest.slide_height_emu,
            "layout_types": sorted(
                {layout.layout_type.value for layout in manifest.layouts}
            ),
        }
    return JSONResponse(
        content={
            "job_id": job_id,
            "title": job.template_filename,
            "purpose": job.purpose,
            "slide_count": job.slide_count,
            "template_dna": template_dna,
            "variants": [
                _variant_payload(job_id, variant, manifest)
                for variant in job.result.variants.values()
            ],
        }
    )


@router.post(
    "/jobs/{job_id}/slides/{variant}/{slide_position}/revise",
    response_model=JobCreatedResponse,
    status_code=202,
)
async def revise_job_slide(
    job_id: str,
    variant: str,
    slide_position: int,
    request: SlideRevisionRequest,
    background_tasks: BackgroundTasks,
    settings: Annotated[Settings, Depends(get_settings)],
) -> JobCreatedResponse:
    job = job_store.get(job_id)
    if job is None or job.result is None:
        raise HTTPException(404, "job not found or not finished")
    code = variant.upper()
    variant_result = job.result.variants.get(code)
    if variant_result is None or variant_result.error is not None:
        raise HTTPException(404, "variant not available")
    manifest = job.result.template_manifest
    if manifest is None:
        raise HTTPException(409, "template analysis is not available")
    if variant_result.presentation_ir is None or not (
        1 <= slide_position <= len(variant_result.presentation_ir.slides)
    ):
        raise HTTPException(404, "slide position is out of range")
    if (
        request.base_revision is not None
        and request.base_revision != variant_result.revision
    ):
        raise HTTPException(
            409,
            f"revision conflict: expected {request.base_revision}, "
            f"current {variant_result.revision}",
        )

    child = job_store.create(
        template_filename=job.template_filename,
        job_type=JobType.REVISE_SLIDE,
        parent_job_id=job_id,
        variant=code,
        slide_position=slide_position,
        revision_request={
            "instructions": request.as_instructions(),
            "base_revision": request.base_revision,
        },
    )
    variant_result.export_state = "STALE"
    job_store.update(job_id, result=job.result)
    _enqueue_job(
        job_id=child.job_id,
        revision=True,
        background_tasks=background_tasks,
        settings=settings,
    )
    return JobCreatedResponse(job_id=child.job_id)


@router.get("/jobs/{job_id}/slides/{variant}/{slide_position}/revisions")
async def list_slide_revisions(
    job_id: str, variant: str, slide_position: int
) -> JSONResponse:
    job = job_store.get(job_id)
    if job is None:
        raise HTTPException(404, "job not found")
    items = revision_store.list(job_id, variant.upper(), slide_position)
    return JSONResponse(
        content={
            "items": [
                {
                    "revision_id": item.revision_id,
                    "revision": item.revision_number,
                    "correction_prompt": item.correction_prompt,
                    "preview_url": (
                        f"/jobs/{job_id}/slides/{variant.upper()}/{slide_position}/"
                        f"revisions/{item.revision_id}/preview"
                        if item.preview_path
                        else None
                    ),
                    "is_active": item.is_active,
                    "created_at": item.created_at.isoformat(),
                }
                for item in items
            ]
        }
    )


@router.get(
    "/jobs/{job_id}/slides/{variant}/{slide_position}/revisions/{revision_id}/preview"
)
async def get_revision_preview(
    job_id: str,
    variant: str,
    slide_position: int,
    revision_id: str,
) -> FileResponse:
    item = revision_store.get(revision_id)
    if (
        item is None
        or item.deck_job_id != job_id
        or item.variant != variant.upper()
        or item.slide_position != slide_position
        or item.preview_path is None
    ):
        raise HTTPException(404, "revision preview not found")
    return FileResponse(item.preview_path, media_type="image/png")


@router.post(
    "/jobs/{job_id}/slides/{variant}/{slide_position}/revisions/{revision_id}/activate",
    response_model=JobCreatedResponse,
    status_code=202,
)
async def activate_slide_revision(
    job_id: str,
    variant: str,
    slide_position: int,
    revision_id: str,
    request: ActivateRevisionRequest,
    background_tasks: BackgroundTasks,
    settings: Annotated[Settings, Depends(get_settings)],
) -> JobCreatedResponse:
    job = job_store.get(job_id)
    code = variant.upper()
    if job is None or job.result is None or code not in job.result.variants:
        raise HTTPException(404, "job or variant not found")
    stored = revision_store.get(revision_id)
    if (
        stored is None
        or stored.deck_job_id != job_id
        or stored.variant != code
        or stored.slide_position != slide_position
    ):
        raise HTTPException(404, "revision not found")
    current = job.result.variants[code]
    if request.base_revision is not None and request.base_revision != current.revision:
        raise HTTPException(
            409,
            f"revision conflict: expected {request.base_revision}, current {current.revision}",
        )
    child = job_store.create(
        template_filename=job.template_filename,
        job_type=JobType.ACTIVATE_REVISION,
        parent_job_id=job_id,
        variant=code,
        slide_position=slide_position,
        revision_request={
            "revision_id": revision_id,
            "base_revision": request.base_revision,
        },
    )
    current.export_state = "STALE"
    job_store.update(job_id, result=job.result)
    _enqueue_job(
        job_id=child.job_id,
        revision=True,
        background_tasks=background_tasks,
        settings=settings,
    )
    return JobCreatedResponse(job_id=child.job_id)


@router.get("/jobs/{job_id}/audit/{variant}")
async def get_job_audit(job_id: str, variant: str) -> JSONResponse:
    job = job_store.get(job_id)
    if job is None or job.result is None:
        raise HTTPException(404, "job not found or not finished")
    variant_result = job.result.variants.get(variant)
    if variant_result is None or variant_result.audit_report is None:
        raise HTTPException(404, f"audit report not available for variant {variant!r}")
    return JSONResponse(content=variant_result.audit_report.model_dump(mode="json"))


@router.get("/jobs/{job_id}/download")
async def download_job_bundle(job_id: str) -> Response:
    job = job_store.get(job_id)
    if job is None or job.result is None:
        raise HTTPException(404, "job not found or not finished")

    buffer = io.BytesIO()
    files_added = 0
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for variant, vr in job.result.variants.items():
            if vr.export_state == "STALE":
                continue
            if vr.pptx_path is not None:
                zf.write(vr.pptx_path, f"variant_{variant}.pptx")
                files_added += 1
            if vr.pdf_path is not None:
                zf.write(vr.pdf_path, f"variant_{variant}.pdf")
                files_added += 1
            if vr.html_path is not None:
                zf.write(vr.html_path, f"variant_{variant}.html")
                files_added += 1
            if vr.audit_report is not None:
                zf.writestr(
                    f"audit_{variant}.json", vr.audit_report.model_dump_json(indent=2)
                )
                files_added += 1

    if files_added == 0:
        raise HTTPException(404, "no files available for this job")

    return Response(
        content=buffer.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="job_{job_id}.zip"'},
    )
